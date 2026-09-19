"""Audit the complete paired 5K comparison and publish inspectable results."""
import argparse
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook
from scipy import linalg

from experiments.guidance_pasted_20260912 import common as c
from experiments.guidance_pasted_20260912.audit import REFS
from . import core as m

OUT = c.WORK/'docs/data/raev2_context_5k_20260914'
REPORT = c.WORK/'docs/RAEV2_CONTEXT_5K_RESULTS_20260914_ZH.md'


def reference_moments():
    with np.load(REFS['raev2']) as d:
        mean = np.asarray(d['mu'],dtype=np.float64)
        cov = np.asarray(d['sigma'],dtype=np.float64)
    eig,vectors = linalg.eigh((cov+cov.T)*.5,check_finite=True)
    assert eig.min() > -1e-8
    root = (vectors*np.sqrt(np.maximum(eig,0)))@vectors.T
    return mean,cov,root


def fid_from_covariance(features,reference):
    """FP64 symmetric covariance formula; size is feature dimension, not N."""
    mean,cov,root = reference
    x = np.asarray(features,dtype=np.float64)
    gen_mean = x.mean(0)
    x = x-gen_mean
    gen_cov = x.T@x/(len(x)-1)
    product = root@gen_cov@root
    eig = linalg.eigvalsh((product+product.T)*.5,check_finite=True)
    assert eig.min() > -1e-7
    return float(np.square(gen_mean-mean).sum()+np.trace(gen_cov)+np.trace(cov)
        -2*np.sqrt(np.maximum(eig,0)).sum())


def audit_arm(arm,reference,noise,labels,request_hash):
    root = m.BASE/arm
    result = c.read(root/'metrics.json')
    summary = c.read(root/'summary.json')
    assert summary['complete'] and result['primary_samples']==m.N
    assert result['request_sha256']==request_hash
    assert c.sha(root/'samples.npz')==result['samples_sha256']==summary['samples_sha256']
    with np.load(root/'samples.npz') as d:
        pixels = d['arr_0']
    assert pixels.dtype==np.uint8 and pixels.shape==(m.N,256,256,3)
    coverage=[]; seconds=0.; batches=[]
    expected_mlp = 99 if arm.startswith('context') else 0
    for record in summary['records']:
        path = Path(record['file'])
        assert c.sha(path)==record['sha256']
        meta = c.read(path.with_suffix('.json'))
        assert meta['sha256']==record['sha256'] and meta['request_sha256']==request_hash
        with np.load(path) as d:
            start = int(d['start']); n = len(d['labels'])
            assert n==min(m.BATCH,m.N-start) and meta['start']==start
            coverage.extend(range(start,start+n))
            np.testing.assert_array_equal(d['labels'],labels[start:start+n])
            np.testing.assert_array_equal(d['arr_0'],pixels[start:start+n])
            assert str(d['noise_sha256'])==c.array_sha(noise[start:start+n])
            assert str(d['request_sha256'])==request_hash and np.isfinite(d['latents']).all()
            assert int(d['full_calls'])==100 and int(d['prefix_calls'])==0
            assert int(d['mlp_calls'])==expected_mlp
            np.testing.assert_array_equal(d['block_calls'],np.full(30,100))
            elapsed = float(d['seconds']); assert elapsed > 0
            seconds += elapsed
            batches.append(dict(arm=arm,start=start,samples=n,seconds=elapsed,
                full_calls=100,prefix_calls=0,mlp_calls=expected_mlp))
    assert coverage==list(range(m.N))
    assert abs(seconds-result['seconds'])<1e-6
    assert result['full_calls_per_output']==100 and result['prefix_calls_at_inference']==0
    assert result['mlp_calls_per_output']==expected_mlp
    feature_paths = list((root/'features').glob('*.features.pt'))
    logit_paths = list((root/'features').glob('*.logits.pt'))
    assert len(feature_paths)==len(logit_paths)==1
    features = torch.load(feature_paths[0],map_location='cpu',weights_only=True).numpy()
    logits = torch.load(logit_paths[0],map_location='cpu',weights_only=True)
    assert features.shape==(m.N,2048) and features.dtype==np.float32 and np.isfinite(features).all()
    assert logits.ndim==2 and len(logits)==m.N and torch.isfinite(logits).all()
    fid = fid_from_covariance(features,reference)
    error = abs(fid-result['fid'])
    assert error<2e-3,(arm,fid,result['fid'],error)
    evaluation = c.read(root/'evaluation_request.json')
    command = evaluation['command']
    assert evaluation['gpu']==2 and evaluation['same_evaluator_device_and_batch_for_both']
    assert command[command.index('--batch-size')+1]=='64'
    assert command[command.index('--device')+1]=='cuda'
    assert evaluation['samples_sha256']==result['samples_sha256']
    official = c.read(root/'official.json')[0]
    csv = pd.read_csv(root/'official.csv').iloc[0]
    assert official==result['metrics']
    assert official['sample_sha256']==result['samples_sha256']
    assert official['fid_reference']=='imagenet_256_fid_stats'
    assert np.isclose(csv.fid,result['fid'],rtol=1e-13,atol=1e-13)
    assert np.isclose(csv.inception_score,result['inception_score'],rtol=1e-13,atol=1e-13)
    row = dict(passed=True,arm=arm,samples=m.N,batches=len(batches),
        coverage_and_paired_inputs_verified=True,pixels_match_raw_batches=True,
        full_calls_per_output=100,extra_prefix_calls=0,mlp_calls_per_output=expected_mlp,
        fid_reported=result['fid'],fid_fp64_from_same_features=fid,absolute_error=error,
        independent_feature_extraction=False,independent_generation_replicate=False,
        samples_sha256=result['samples_sha256'],features_sha256=c.sha(feature_paths[0]),
        logits_sha256=c.sha(logit_paths[0]),reference_sha256=c.sha(REFS['raev2']),
        evaluator_commit=official['evaluator_commit'])
    c.atomic(root/'audit.json',row)
    print(arm,'arithmetic / sample audit passed; FID absolute error',error,flush=True)
    return row,batches,pixels[:4].copy()


def write_tables(frames):
    for name,frame in frames.items():
        frame.to_csv(OUT/(name+'.csv'),index=False)
    with pd.ExcelWriter(OUT/'source_data.xlsx',engine='openpyxl') as writer:
        for name,frame in frames.items():
            frame.to_excel(writer,sheet_name=name,index=False)
            writer.book[name].freeze_panes='A2'
    book = load_workbook(OUT/'source_data.xlsx',data_only=True,read_only=True)
    for name,frame in frames.items():
        assert book[name].max_row==len(frame)+1
        values = book[name].iter_rows(values_only=True)
        assert list(next(values))==list(frame.columns)
        for expected,actual in zip(frame.itertuples(index=False,name=None),values):
            for x,y in zip(expected,actual):
                if isinstance(x,(float,int,np.number)) and not isinstance(x,(bool,np.bool_)):
                    assert np.isclose(x,y,rtol=1e-12,atol=1e-12)
                else:
                    assert x==y
    book.close()


def build():
    m.configure(); torch.set_num_threads(4)
    status = c.read(m.ROOT/'status.json')
    assert status['phase']=='complete' and status['quality_images']==10000
    request = m.verify(); rh = c.sha(m.BASE/'request.json')
    assert request['arms']==list(m.ARMS) and request['samples']==5000
    preflight = c.read(m.ROOT/'preflight.json'); assert preflight['passed']
    noise,labels = m.bank()
    unique,counts = np.unique(labels,return_counts=True)
    assert np.array_equal(unique,np.arange(1000)) and np.all(counts==5)
    reference = reference_moments()
    rows = c.read(m.BASE/'results.json')
    assert [r['arm'] for r in rows]==list(m.ARMS)
    audits=[]; batch_rows=[]; first_pixels=[]
    for arm in m.ARMS:
        audit,batches,pixels = audit_arm(arm,reference,noise,labels,rh)
        audits.append(audit);batch_rows.extend(batches);first_pixels.append(pixels)
    assert audits[0]['evaluator_commit']==audits[1]['evaluator_commit']
    benchmark = c.read(m.ROOT/'benchmark.json')
    assert benchmark['complete'] and benchmark['batch']==m.BATCH
    assert benchmark['warmups']==1 and benchmark['repeats']==3
    assert benchmark['full_calls']==100 and benchmark['extra_prefix']==0
    timing=[]
    for arm in m.ARMS:
        values = benchmark['seconds'][arm]
        assert len(values)==3 and abs(np.median(values)-benchmark['medians'][arm])<1e-12
        timing.extend(dict(arm=arm,repeat=i+1,batch=m.BATCH,seconds=t) for i,t in enumerate(values))
    quality = pd.DataFrame([dict(arm=r['arm'],alpha=request['alpha'][r['arm']],
        official_scale=1+request['alpha'][r['arm']],samples=m.N,classes=1000,images_per_class=5,
        fid=r['fid'],inception_score=r['inception_score'],full_calls=100,extra_prefix_calls=0,
        mlp_calls=r['mlp_calls_per_output'],benchmark_seconds_per_batch=benchmark['medians'][r['arm']],
        batch=m.BATCH,samples_sha256=r['samples_sha256']) for r in rows])
    native,mlp = quality.iloc[0],quality.iloc[1]
    delta = float(mlp.fid-native.fid)
    relative = float(mlp.fid/native.fid-1)
    timing_change = float(mlp.benchmark_seconds_per_batch/native.benchmark_seconds_per_batch-1)
    assert abs(timing_change-benchmark['relative_mlp_change'])<1e-12
    comparison = pd.DataFrame([dict(native_fid=native.fid,mlp_20k_fid=mlp.fid,delta_mlp_minus_native=delta,
        relative_fid_change=relative,inception_score_change=float(mlp.inception_score-native.inception_score),
        relative_sampling_time_change=timing_change,coefficients_differ=True,paired_seed=m.SEED)])
    OUT.mkdir(parents=True,exist_ok=True)
    write_tables({'quality':quality,'comparison':comparison,'timing_repeats':pd.DataFrame(timing),
        'sampling_batches':pd.DataFrame(batch_rows),'class_counts':pd.DataFrame({'class':unique,'samples_per_arm':counts})})
    plt.rcParams.update({'font.size':11,'axes.spines.top':False,'axes.spines.right':False,
        'svg.fonttype':'none','pdf.fonttype':42})
    fig,axes = plt.subplots(1,2,figsize=(8.3,4.1),layout='constrained')
    names=['Native IG\nalpha 0.78','20K MLP IG\nalpha 0.39']
    for ax,column,title in zip(axes,['fid','inception_score'],['FID (lower is better)','IS (higher is better)']):
        bars = ax.bar(names,quality[column],width=.56,color=['#39718A','#BB8241'])
        ax.bar_label(bars,fmt='%.3f',padding=4,fontsize=10)
        ax.set_ylim(0,float(quality[column].max())*1.15)
        ax.set_title(title,fontsize=11)
        ax.grid(axis='y',alpha=.18); ax.set_axisbelow(True)
    fig.suptitle('RAEv2: 5,000 images per configuration',fontsize=13)
    for ext in ('png','svg','pdf'):
        fig.savefig(OUT/f'quality_comparison.{ext}',dpi=170)
    plt.close(fig)
    fig,axes = plt.subplots(2,4,figsize=(9.3,4.8),layout='constrained')
    for row,pixels in enumerate(first_pixels):
        for col,pixel in enumerate(pixels):
            ax=axes[row,col];ax.imshow(pixel);ax.set_xticks([]);ax.set_yticks([])
            for spine in ax.spines.values():spine.set_visible(False)
            if row==0:ax.set_title(f'Index {col}; class {labels[col]}',fontsize=9)
            if col==0:ax.set_ylabel(names[row],rotation=0,ha='right',va='center',fontsize=9)
    fig.suptitle('Same noise and class: first four samples in fixed order',fontsize=12)
    fig.savefig(OUT/'paired_first4.png',dpi=160);plt.close(fig)
    shown = quality[['alpha','samples','fid','inception_score','benchmark_seconds_per_batch']].copy()
    shown.insert(0,'设置',['原生IG','20K Context MLP IG'])
    shown.rename(columns={'alpha':'额外系数α','samples':'图数','fid':'FID↓','inception_score':'IS↑',
        'benchmark_seconds_per_batch':'每16图耗时/秒'},inplace=True)
    direction = '降低' if delta<0 else '升高'
    conclusion = ('此次5K比较中，20K MLP设置的FID更低。' if delta<0 else
        '此次5K比较中，20K MLP设置的FID没有优于原生IG。')
    paragraphs = [
        '**RAEv2：原生IG与20K MLP的5K比较已完成。** '+conclusion,
        f'每臂全新生成5000张，覆盖1000类、每类5张；共享种子{m.SEED}产生的初始噪声和标签顺序。'
        f'20K MLP相对原生IG的FID变化为{delta:+.4f}（{direction}{abs(relative):.2%}），'
        f'IS变化为{float(mlp.inception_score-native.inception_score):+.4f}。',
        shown.to_markdown(index=False,floatfmt=('', '.2f', '.0f', '.4f', '.4f', '.4f')),
        '原生IG使用额外α=.78，即官方scale=1.78；20K MLP使用额外α=.39，即scale=1.39。'
        '两者分别取此前各自两档400图测试中FID最低的设置，在本轮生成前固定。此次比较的是两个完整配置；'
        '因系数不同，差异不能单独归因于MLP结构、弱头训练或20K续训。5K采用新噪声，未复用旧400图质量样本。',
        '![FID与IS](data/raev2_context_5k_20260914/quality_comparison.png)',
        '两者均使用同一冻结RAEv2主干、解码器、Euler100、时间shift8和活动窗口[.1,1]；'
        '采样batch=16，最后一批8，bf16 autocast。每张图100次完整主干调用，0额外prefix调用；'
        '30个Transformer block的实际调用计数全部核对。MLP在99个活动步复用编码器第8层特征，'
        '新增5,635,744个参数；原生弱头仍随主干计算。本次没有新增训练。',
        f'耗时在同一张GPU1上测量，包含100步采样及解码；每臂预热1次，再交错测量3次，表中取中位数。'
        f'MLP设置耗时变化为{timing_change:+.2%}。该短基准用于记录实际开销，不能据此宣称严格零成本。',
        '所有10000张质量输出均进入评价，无筛图；两臂使用同一GPU2、FP32 Inception和batch64，'
        '同一nanogen-evals版本与imagenet_256_fid_stats。FID采用相同参考统计；'
        '用保存的5000×2048特征按FP64对称协方差公式复算，逐臂与原评价器结果核对。'
        '复算只是同一批特征的算术核验，不是独立特征提取或独立生成重复。',
        '本轮单一生成bank和单一训练种子没有提供独立重复或置信区间；'
        '不将5K绝对FID与此前400图或论文50K结果直接比较，也不据此验证共同误差分布的机制。',
        '![配对前四张](data/raev2_context_5k_20260914/paired_first4.png)',
        '上图仅展示固定顺序前4张；不是人工挑选的样例。新质量采样之前各重放4张旧样本，'
        'latent及像素逐位一致（旧重放batch4，两臂正式5K统一batch16）。'
        '另有耗时测试8批×16图，共128次生成，未计入FID/IS；本次共10000张质量样本、8张重放和128次基准生成。',
        '[冻结协议](RAEV2_CONTEXT_5K_PROTOCOL_20260914_ZH.md) · '
        '[源数据工作簿](data/raev2_context_5k_20260914/source_data.xlsx) · '
        '[质量CSV](data/raev2_context_5k_20260914/quality.csv) · '
        '[核验记录](data/raev2_context_5k_20260914/verification.json)。',
        f'[原生IG完整5K样本]({m.BASE/"native_base/samples.npz"}) · '
        f'[20K MLP完整5K样本]({m.BASE/"context20k_half/samples.npz"})。',
        '此次明确5K请求已执行完，两臂均完成全部5000张。此前400图自动扩展门槛失败记录保留；'
        '旧大队列未恢复，未自动开展新训练或系数搜索。']
    REPORT.write_text('\n\n'.join(paragraphs)+'\n')
    workers = c.read(m.ROOT/'workers.json'); processes=[]
    for pid in [workers['parent']]+[r['pid'] for r in workers['children']]:
        path = Path('/proc')/str(pid)/'cmdline'
        command = path.read_bytes().decode(errors='replace') if path.exists() else ''
        assert 'experiments.raev2_context_5k_20260914.run' not in command
        assert 'experiments.raev2_context_5k_20260914.core' not in command
        processes.append(dict(pid=pid,experiment_ended=True))
    historical = c.read(c.EXPS/'readout_followup_completion_20260913.json')['old_queue_stop_markers']
    for name,digest in historical.items():
        assert c.sha(c.EXPS/name/'STOP_AFTER_CURRENT')==digest
    old_screen_stop = m.tr.ROOT/'STOP_AFTER_SCREEN'
    assert old_screen_stop.read_text()=='Neither fixed 20K strength passed the paired 400-image gate; no 1K confirmation or strength search.\n'
    verification = dict(passed=True,new_quality_images=10000,samples_per_arm=5000,
        classes=1000,images_per_class=5,old_replay_images=8,benchmark_generated_images=128,
        source_request_sha256=rh,final_head_sha256=c.sha(m.tr.TRAIN/'head.pt'),
        source_workbook_reconciled=True,paired_noise_and_labels=True,coefficients_differ=True,
        audits=audits,process_checks=processes,old_queue_stop_markers=historical,
        historical_400_gate_stop_sha256=c.sha(old_screen_stop),
        additional_training=False,visual_review_pending=True,goal_complete=False)
    c.atomic(OUT/'verification.json',verification)
    update_manifest()
    print('Report prepared; visual inspection remains:',REPORT,flush=True)


def update_manifest():
    c.atomic(OUT/'manifest.json',dict(report_sha256=c.sha(REPORT),generator_sha256=c.sha(Path(__file__)),
        files={p.name:c.sha(p) for p in OUT.iterdir() if p.is_file() and p.name!='manifest.json'}))


def finalize_visual_review():
    manifest = c.read(OUT/'manifest.json')
    assert manifest['report_sha256']==c.sha(REPORT)
    assert manifest['generator_sha256']==c.sha(Path(__file__))
    for name,digest in manifest['files'].items():
        assert c.sha(OUT/name)==digest,name
    verification = c.read(OUT/'verification.json'); assert verification['passed']
    verification['visual_review_pending']=False
    verification['visuals_reviewed']=['quality_comparison.png','paired_first4.png']
    c.atomic(OUT/'verification.json',verification);update_manifest()
    c.atomic(m.ROOT/'completion_verification.json',dict(passed=True,user_request_completed=True,
        request='原生IG与20K MLP各5000张配对比较',quality_images=10000,old_replay_images=8,
        benchmark_generated_images=128,additional_training=False,
        report=str(REPORT),report_sha256=c.sha(REPORT),manifest_sha256=c.sha(OUT/'manifest.json'),
        final_head_sha256=verification['final_head_sha256'],goal_complete=False))
    print('5K comparison finalized:',REPORT,flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--finalize-visual-review',action='store_true')
    args=parser.parse_args()
    if args.finalize_visual_review:finalize_visual_review()
    else:build()
