"""CPU proofs, counterexamples, and existing-trace audit for AG gain schedules.

Run: PYTHONPATH=/tmp/eqvae_ag_audit_dependencies python \
     experiments/audit_ag_ctrl_decay_20260911.py
NumPy, SciPy, pandas, matplotlib, torch, and openpyxl are required. No model
checkpoint, image sampler, training job, or GPU is loaded. The image trace is
an EXISTING SiT CFG-Ctrl adaptation, not a new AG image experiment.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from scipy.integrate import solve_ivp

from ag_error_cancellation_20260911 import apply_ag, calibrate_risk_gamma, hierarchy_gamma


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/data/ag_ctrl_decay_20260911"
READINGS = ROOT / "readings/ag_ctrl_decay_20260911"
TRACE = ROOT / "docs/data/sit_fsg_ctrl_hypothesis_20260911/ctrl_trace_rows.csv"
CTRL = ROOT / "readings/fsg_ctrl_mp_comparison_20260911/code/cfg_ctrl/pipeline/common_cfg_ctrl.py"
BLUE, GRAY, ORANGE = "#245879", "#555555", "#9b5928"


def scalar(value):
    return float(value.detach().cpu()) if isinstance(value, torch.Tensor) else float(value)


def save_frame(name, rows):
    frame = rows if isinstance(rows, pd.DataFrame) else pd.DataFrame(rows)
    frame.to_csv(OUT / f"{name}.csv", index=False)
    return frame


def replay_ctrl():
    spec = importlib.util.spec_from_file_location("ag_audit_author_ctrl", CTRL)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    instance = module.CFGCtrlMixin()
    state = module.CFGCtrlState()
    params = module.CFGCtrlParams(smc_cfg_enable=True, smc_cfg_lambda=5.0, smc_cfg_K=0.2)
    rows = []
    for k in range(12):
        result = instance._cfg_ctrl_apply(
            noise_pred_posi=torch.tensor([0.05], dtype=torch.float64),
            noise_pred_nega=torch.zeros(1, dtype=torch.float64),
            cfg_scale=1.0, progress_id=k, params=params, state=state,
        )
        rows.append(dict(step=k, raw_gap=0.05, modified_gap=scalar(result[0])))
    expected = np.tile([-0.15, 0.25], 6)
    np.testing.assert_allclose([r["modified_gap"] for r in rows], expected, atol=1e-14)
    return save_frame("author_ctrl_cycle", rows)


def trace_scalar_audit():
    frame = pd.read_csv(TRACE)
    frame = frame.loc[frame.method == "smc_high"].copy()
    assert len(frame) == 200 * 48
    assert frame.groupby("k").size().eq(200).all()
    assert not frame.duplicated(["start", "index", "k"]).any()
    g2 = frame.gap_rms.to_numpy() ** 2
    m2 = frame.modified_rms.to_numpy() ** 2
    c2 = frame.correction_rms.to_numpy() ** 2
    dot_per_dimension = (g2 + m2 - c2) / 2
    alpha = dot_per_dimension / g2
    perp2 = m2 - dot_per_dimension ** 2 / g2
    assert float(np.min(perp2)) >= -1e-8
    frame["best_scalar_alpha"] = alpha
    frame["orthogonal_fraction"] = np.sqrt(np.maximum(perp2, 0) / m2)
    frame["sample_id"] = frame.start + frame["index"]
    save_frame("ctrl_scalar_projection_rows", frame)
    rows = []
    for k, part in frame.groupby("k"):
        rows.append(dict(
            k=int(k), progress=float(k / 64), n_samples=len(part),
            raw_gap_rms_mean=float(part.gap_rms.mean()),
            modified_gap_rms_mean=float(part.modified_rms.mean()),
            orthogonal_fraction_median=float(part.orthogonal_fraction.median()),
            orthogonal_fraction_p10=float(part.orthogonal_fraction.quantile(.1)),
            orthogonal_fraction_p90=float(part.orthogonal_fraction.quantile(.9)),
            best_scalar_alpha_median=float(part.best_scalar_alpha.median()),
        ))
    return save_frame("ctrl_scalar_projection_by_step", rows)


def gaussian_audit():
    rows, loss_rows, checks = [], [], []
    sigmas = np.unique(np.r_[np.geomspace(.001, 30, 241), .01, .1, 1, 3, 10])[::-1]
    for sigma in sigmas:
        u, vt, vs, vw = sigma ** 2, 1.0, 2.0, 4.0
        at, ass, aw = vt / (vt + u), vs / (vs + u), vw / (vw + u)
        gap = ass - aw
        gamma = .5 * (4 + u) / (1 + u)
        gap_rms = abs(gap) * np.sqrt(vt + u)
        correction = gamma * gap_rms
        score_precision = (1 + gamma) / (vs + u) - gamma / (vw + u)
        ode_residual = -1.5 / (1 + u) ** 2 + (1 / (vs + u) - 1 / (vw + u)) * gamma * (1 + gamma)
        checks.append(abs(ass + gamma * gap - at))
        checks.append(abs(score_precision - 1 / (vt + u)))
        checks.append(abs(ode_residual))
        rows.append(dict(
            sigma=sigma, noise_variance=u, target_variance=vt, strong_variance=vs,
            weak_variance=vw, gamma=gamma, edm2_guidance=1 + gamma,
            raw_clean_gap_rms=gap_rms, clean_correction_rms=correction,
            ode_drift_correction_rms=correction / sigma,
            score_correction_rms=correction / u,
            tilted_marginal_variance=1 / score_precision,
            target_marginal_variance=vt + u,
        ))
    for vt in [1., 2., 2.5]:
        for sigma in [10., 3., 1., .1, .01]:
            u, vs, vw = sigma ** 2, 2., 4.
            clean = torch.tensor([[-np.sqrt(vt)], [-np.sqrt(vt)], [np.sqrt(vt)], [np.sqrt(vt)]], dtype=torch.float64)
            noise = torch.tensor([[-sigma], [sigma], [-sigma], [sigma]], dtype=torch.float64)
            noisy = clean + noise
            strong = vs / (vs + u) * noisy
            weak = vw / (vw + u) * noisy
            result = calibrate_risk_gamma(clean, strong, weak)
            analytic_gamma = (vt / (vt + u) - vs / (vs + u)) / (vs / (vs + u) - vw / (vw + u))
            # Near sigma=0 the gap is tiny, so dividing a roundoff-sized
            # covariance by B is ill-conditioned. Check the resulting
            # prediction error as well as an absolute coefficient tolerance.
            np.testing.assert_allclose(scalar(result["gamma"]), analytic_gamma, atol=2e-8)
            assert abs(scalar(result["gamma"]) - analytic_gamma) * np.sqrt(scalar(result["B"])) < 1e-13
            np.testing.assert_allclose(scalar(result["A"]), scalar((result["weak_mse"] - result["strong_mse"] - result["B"]) / 2), atol=2e-14)
            assert abs(scalar(result["orthogonality"])) < 1e-14
            assert scalar(result["strong_mse"]) < scalar(result["weak_mse"])
            loss_rows.append(dict(target_variance=vt, sigma=sigma,
                                  **{k:scalar(v) for k,v in result.items()}))
    assert max(checks) < 2e-14, max(checks)
    return save_frame("gaussian_gain_and_correction", rows), save_frame("risk_calibration_and_nonidentifiability", loss_rows), max(checks)


def hierarchy_audit():
    rows, max_error = [], 0.
    for kind in ["constant_ratio", "growing_ratio", "falling_ratio"]:
        for tau in np.linspace(0, 1, 101):
            b = {"constant_ratio": 2., "growing_ratio": .5 * 4 ** tau,
                 "falling_ratio": 2 * 4 ** (-tau)}[kind]
            kappa = 1 + b
            truth = torch.tensor([[.3, -.7, .2]], dtype=torch.float64)
            e = (1 - .8 * tau) * torch.tensor([[.7, -.3, .1]], dtype=torch.float64)
            predictions = [truth + kappa ** j * e for j in range(3)]
            result = hierarchy_gamma(*predictions)
            estimate = apply_ag(predictions[0], predictions[1], result["gamma"])
            error = scalar((estimate - truth).norm())
            max_error = max(max_error, error)
            np.testing.assert_allclose(scalar(result["gamma"][0]), 1 / b, atol=3e-14)
            transformed = hierarchy_gamma(*[-3.7 * d + 1.2 for d in predictions])
            np.testing.assert_allclose(transformed["gamma"], result["gamma"], atol=3e-14)
            rows.append(dict(
                family=kind, sampling_progress=tau, kappa=kappa, gamma=scalar(result["gamma"][0]),
                strong_error_rms=scalar(e.square().mean().sqrt()),
                guided_error_l2=error, relative_equality_residual=scalar(result["residual_relative"][0]),
            ))
    predictions = [torch.tensor([[x]], dtype=torch.float64) for x in [.1, 1.1, 3.1]]
    bad = hierarchy_gamma(*predictions)
    bad_out = apply_ag(predictions[0], predictions[1], bad["gamma"])
    np.testing.assert_allclose(scalar(bad["gamma"][0]), 1., atol=1e-14)
    np.testing.assert_allclose(scalar(bad_out[0, 0]), -.9, atol=1e-14)
    same = torch.ones((2, 3), dtype=torch.float64)
    degenerate = hierarchy_gamma(same, same, same)
    assert not degenerate["identified"].any()
    assert torch.isfinite(degenerate["gamma"]).all()
    np.testing.assert_array_equal(apply_ag(same, same, degenerate["gamma"]), same)
    # Nonparallel mismatch: regression optimum differs from 1 / relative bias.
    mismatch = calibrate_risk_gamma(torch.zeros((1,2), dtype=torch.float64),
                                   torch.tensor([[1.,0.]], dtype=torch.float64),
                                   torch.tensor([[3.,-1.]], dtype=torch.float64))
    np.testing.assert_allclose(scalar(mismatch["gamma"]), .4, atol=1e-14)
    assert max_error < 2e-14
    counterexample = dict(
        true_target=0., strong=.1, weak=1.1, weaker=3.1,
        gamma=scalar(bad["gamma"][0]), extrapolated=scalar(bad_out[0,0]),
        strong_squared_error=.01, extrapolated_squared_error=.81,
        equality_residual=scalar(bad["residual_relative"][0]),
        warning="Exact hierarchy equality does not identify the true limit; common bias can make extrapolation worse.",
    )
    return save_frame("hierarchy_schedules", rows), counterexample, max_error


def gaussian_transport():
    u_hi, u_lo, vt, vs, vw = 100., .0001, 1., 2., 4.
    aa = (vs + u_lo) / (vs + u_hi)
    bb = (vw + u_lo) / (vw + u_hi)
    constant_fit = (np.log((vt + u_lo) / (vt + u_hi)) - np.log(aa)) / np.log(aa / bb)
    rows, endpoint_rows = [], []
    for name, gamma_function in [
        ("exact_path", lambda u:.5 * (4 + u)/(1 + u)),
        ("constant_0.5", lambda u:.5),
        ("constant_2", lambda u:2.),
        ("constant_endpoint_fitted", lambda u:constant_fit),
    ]:
        def rhs(u, log_variance):
            gamma = gamma_function(u)
            return [(1 + gamma)/(vs + u) - gamma/(vw + u)]
        grid = np.geomspace(u_hi, u_lo, 121)
        sol = solve_ivp(rhs, (u_hi,u_lo), [np.log(vt + u_hi)], t_eval=grid,
                        rtol=2e-11, atol=2e-12, method="DOP853")
        assert sol.success
        variances = np.exp(sol.y[0])
        for u, variance in zip(sol.t, variances):
            rows.append(dict(method=name, noise_variance=u, sigma=np.sqrt(u),
                             gamma=gamma_function(u), variance=variance,
                             target_variance=vt+u, variance_ratio=variance/(vt+u)))
        endpoint_rows.append(dict(method=name, endpoint_variance=variances[-1],
                                  target_endpoint_variance=vt+u_lo,
                                  endpoint_wasserstein2=(np.sqrt(variances[-1])-np.sqrt(vt+u_lo))**2,
                                  max_path_variance_ratio_error=float(np.max(np.abs(variances/(vt+sol.t)-1)))))
        if name == "exact_path":
            np.testing.assert_allclose(variances, vt + grid, rtol=2e-9)
        if name == "constant_endpoint_fitted":
            np.testing.assert_allclose(variances[-1], vt + u_lo, rtol=2e-9)
    return save_frame("gaussian_transport_path", rows), save_frame("gaussian_transport_endpoints", endpoint_rows), constant_fit


def figures(gaussian, hierarchy, ctrl):
    plt.rcParams.update({"font.family":"DejaVu Sans", "font.size":10,
                         "axes.spines.top":False, "axes.spines.right":False,
                         "figure.facecolor":"white", "savefig.facecolor":"white"})
    fig, axes = plt.subplots(1, 3, figsize=(13.5, 3.8), constrained_layout=True)
    axes[0].semilogx(gaussian.sigma, gaussian.gamma, color=BLUE, linewidth=2)
    axes[0].set(xlim=(30,.001), xlabel=r"Noise $\sigma$ (sampling: left to right)",
                ylabel=r"Extra AG coefficient $\gamma$", title="Exact Gaussian denoiser correction")
    axes[1].loglog(gaussian.sigma, gaussian.raw_clean_gap_rms, color=GRAY, label="Raw strong-weak gap")
    axes[1].loglog(gaussian.sigma, gaussian.clean_correction_rms, color=BLUE, label="Applied clean correction")
    axes[1].set(xlim=(30,.001), xlabel=r"Noise $\sigma$ (sampling: left to right)",
                ylabel="RMS in clean-prediction units", title="Gap and applied correction")
    axes[1].legend(frameon=False, fontsize=8)
    for family, color, label in [("constant_ratio", GRAY, "Constant error ratio"),
                                 ("growing_ratio", BLUE, "Growing error ratio"),
                                 ("falling_ratio", ORANGE, "Falling error ratio")]:
        part = hierarchy.loc[hierarchy.family == family]
        axes[2].plot(part.sampling_progress, part.gamma, color=color, label=label)
    axes[2].set(xlabel="Sampling progress (constructed family)", ylabel=r"Estimated $\gamma$",
                title="Compatible error hierarchies")
    axes[2].legend(frameon=False, fontsize=8)
    for ax in axes: ax.grid(alpha=.13)
    fig.savefig(OUT / "ag_gain_and_correction.png", dpi=180)
    fig.savefig(OUT / "ag_gain_and_correction.pdf")
    plt.close(fig)
    fig, axes = plt.subplots(1,2,figsize=(9,3.4), constrained_layout=True)
    axes[0].plot(ctrl.progress, ctrl.raw_gap_rms_mean, color=GRAY, label="Raw gap")
    axes[0].plot(ctrl.progress, ctrl.modified_gap_rms_mean, color=BLUE, label="CTRL modified gap")
    axes[0].set(xlabel="SiT sampling progress k/64", ylabel="Mean per-image velocity RMS",
                title="Existing 200-image SiT adaptation")
    axes[0].legend(frameon=False)
    axes[1].fill_between(ctrl.progress, ctrl.orthogonal_fraction_p10,
                         ctrl.orthogonal_fraction_p90, color=BLUE, alpha=.13, label="10–90% across images")
    axes[1].plot(ctrl.progress, ctrl.orthogonal_fraction_median, color=BLUE, label="Median")
    axes[1].set(xlabel="SiT sampling progress k/64", ylabel=r"$\|m-\alpha_*g\|/\|m\|$",
                ylim=(0,1.02), title="Residual after best scalar fit")
    axes[1].legend(frameon=False, fontsize=8, loc="lower right")
    for ax in axes: ax.grid(alpha=.13)
    fig.savefig(OUT / "ctrl_scalar_audit.png", dpi=180)
    fig.savefig(OUT / "ctrl_scalar_audit.pdf")
    plt.close(fig)


def workbook(frames):
    book = Workbook()
    readme = book.active
    readme.title = "Readme"
    readme.append(["AG gain / error-cancellation audit", "Scope"])
    readme.append(["Evidence", "CPU exact models plus reanalysis of an existing SiT CFG-Ctrl trace. No new AG image-quality result."])
    readme.append(["Gamma convention", "Strong + gamma*(Strong-Weak); EDM2 guidance = 1+gamma."])
    readme.append(["Time", "Gaussian sigma decreases during sampling. SiT progress and synthetic hierarchy progress increase."])
    readme.append(["Units", "Gaussian correction RMS is clean-denoiser output; CTRL trace RMS is velocity. Do not compare magnitudes across the two."])
    readme.append(["Gaussian moments", "Four symmetric clean/noise pairs integrate all quadratic denoising losses exactly; they are not image observations."])
    readme.append(["Risk calibration", "Uses true noised-data pairs at a fixed noise level; not a terminal FID objective."])
    readme.append(["Counterexample", "Three-model equality can be exact while squared error increases from 0.01 to 0.81."])
    for title, frame, source in frames:
        sheet = book.create_sheet(title)
        sheet.append(list(frame.columns))
        sheet.cell(1,1).comment = Comment(source, "Source")
        for row in frame.itertuples(index=False, name=None):
            sheet.append([v.item() if isinstance(v, np.generic) else v for v in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    sources = book.create_sheet("Sources")
    sources.append(["Source", "Title", "Version/date", "URL or local path", "Used for"])
    for entry in json.loads((READINGS / "source_manifest.json").read_text()):
        sources.append([entry["key"],entry.get("title",""),"Accessed 2026-09-11; version in URL",entry["url"],"Paper context; see cited report"])
    sources.append(["CTRL code","CFG-Ctrl common implementation","490a628fb0999b9269541e89820e1c26b71931bc",
                    "https://github.com/hanyang-21/CFG-Ctrl/blob/490a628fb0999b9269541e89820e1c26b71931bc/pipeline/common_cfg_ctrl.py","Actual function replay"])
    sources.append(["Existing SiT trace","smc_high, 200 images, active steps 0–47","2026-09-11",str(TRACE),"Scalar projection reanalysis"])
    sources.append(["Analytic derivations","Exact Gaussian and geometric-error examples","2026-09-11",str(Path(__file__).resolve()),"All constructed model data"])
    for sheet in book:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid",fgColor="333333")
        for cells in sheet.columns:
            length = max(len(str(c.value or "")) for c in list(cells)[:100])
            sheet.column_dimensions[cells[0].column_letter].width = min(65,max(14,length+2))
    path = OUT / "ag_ctrl_decay_support.xlsx"
    book.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    assert set(check.sheetnames) == {"Readme","Sources"} | {v[0] for v in frames}
    for title, frame, _ in frames:
        assert check[title].max_row == len(frame)+1
    check.close()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    cycle = replay_ctrl()
    ctrl = trace_scalar_audit()
    gaussian, losses, gaussian_error = gaussian_audit()
    hierarchy, counterexample, hierarchy_error = hierarchy_audit()
    transport, endpoints, constant_fit = gaussian_transport()
    figures(gaussian,hierarchy,ctrl)
    workbook([
        ("Gaussian_gain",gaussian,"Original analytic derivation; target N(0,1), strong N(0,2), weak N(0,4). See report and audit script."),
        ("Hierarchy_schedules",hierarchy,"Constructed compatible-error sequences; CPU verification, not image metrics."),
        ("Risk_and_ambiguity",losses,"Exact quadratic moments of three alternative target Gaussians with the same model pair."),
        ("CTRL_by_step",ctrl,str(TRACE)+"; only smc_high; identity reconstructed from three per-sample RMS norms."),
        ("CTRL_rows",pd.read_csv(OUT / "ctrl_scalar_projection_rows.csv"),str(TRACE)+"; all 200*48 smc_high rows and derived projection quantities."),
        ("CTRL_cycle",cycle,str(CTRL)+"; actual author function on a constant scalar gap."),
        ("Common_bias_failure",pd.DataFrame([counterexample]),"Constructed three-model common-bias counterexample; see audit.json and report."),
        ("Transport_path",transport,"CPU integration of analytic Gaussian probability-flow variance; initialized from true noisy target."),
        ("Transport_endpoints",endpoints,"Shows a calibrated constant can also match the endpoint in 1D; does not establish schedule superiority."),
    ])
    report = dict(
        scope="CPU mathematics and existing SiT trace reanalysis; no new AG image generation or FID",
        checks="passed",
        gaussian_max_identity_error=gaussian_error,
        hierarchy_max_target_l2_error=hierarchy_error,
        common_bias_counterexample=counterexample,
        gaussian_constant_endpoint_fit=constant_fit,
        ctrl_examples=ctrl.loc[ctrl.k.isin([0,16,32,47])].to_dict(orient="records"),
        ctrl_actual_source_sha256=hashlib.sha256(CTRL.read_bytes()).hexdigest(),
        ctrl_input_trace_sha256=hashlib.sha256(TRACE.read_bytes()).hexdigest(),
        versions=dict(python=sys.version.split()[0],numpy=np.__version__,torch=torch.__version__,pandas=pd.__version__),
    )
    (OUT/"audit.json").write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n")
    print(json.dumps(report,indent=2,ensure_ascii=False))


if __name__ == "__main__":
    main()
