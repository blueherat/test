"""Reconcile the 20K loss curves, paired image results, and their source tables."""
import argparse
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook

from experiments.guidance_pasted_20260912 import common as c
from experiments.guidance_distribution_20260912 import audit,mixture
from . import sample as m,train as tr

OUT = c.WORK/'docs/data/raev2_context_20k_20260914'
REPORT = c.WORK/'docs/RAEV2_CONTEXT_20K_RESULTS_20260914_ZH.md'


def root_for(stage,arm):
    if stage == m.SCREEN and arm in m.CONTROL_ARMS:
        return m.OLD_SCREEN/arm
    return m.ROOT/'raev2'/stage/arm


def build(reviewed=False):
    m.configure()
    torch.set_num_threads(4)
    state = c.read(m.ROOT/'status.json')
    assert state['phase'] == 'complete'
    old_request = tr.prepare()
    assert c.sha(old_request) == c.read(tr.TRAIN/'summary.json')['request_sha256']
    train = c.read(tr.TRAIN/'summary.json')
    assert train['steps'] == 20000 and train['replay_exact']
    assert c.sha(tr.TRAIN/'head.pt') == train['head_sha256']
    assert c.read(tr.TRAIN/'replay_3000.json')['passed']
    assert c.read(tr.TRAIN/'old_validation_replay.json')['passed']
    assert c.read(m.ROOT/'preflight.json')['passed']
    OUT.mkdir(parents=True,exist_ok=True)
    history = pd.DataFrame(train['history'])
    assert history.step.tolist() == list(range(1,20001))
    online_blocks = history.assign(block=(history.step-1)//100).groupby('block').agg(
        last_step=('step','max'),context_mean=('context','mean'),observations=('context','count')).reset_index(drop=True)
    assert online_blocks.observations.eq(100).all()
    validation = pd.DataFrame(train['validation'])
    assert validation.step.tolist() == list(range(3000,20001,1000))
    pairs = []
    for p in sorted((tr.TRAIN/'fixed_validation').glob('step_*.json')):
        record = c.read(p)
        for split,entry in record['splits'].items():
            assert len(entry['values']) == 1000
            assert abs(np.mean(entry['values'])-entry['mean']) < 1e-12
            wanted = float(validation.loc[validation.step==record['step'],split].iloc[0])
            assert abs(wanted-entry['mean']) < 1e-12
            pairs.extend(dict(step=record['step'],split=split,case=i,mse=value) for i,value in enumerate(entry['values']))
    pair_loss = pd.DataFrame(pairs)
    quality, audits = [], []
    stages = [m.SCREEN] + ([m.CONFIRM] if state['confirmed'] else [])
    for stage in stages:
        request = m.verify(stage)
        new = c.read(m.ROOT/'raev2'/stage/'results.json')
        assert [r['arm'] for r in new] == request['arms']
        rows = ([c.read(m.OLD_SCREEN/a/'metrics.json') for a in m.CONTROL_ARMS] if stage == m.SCREEN else []) + new
        for row in rows:
            arm = row['arm']
            root = root_for(stage,arm)
            assert c.sha(root/'samples.npz') == row['samples_sha256']
            reused = stage == m.SCREEN and arm in m.CONTROL_ARMS
            if reused:
                old_audit = c.read(root/'audit.json')
                assert old_audit['passed'] and old_audit['samples_sha256'] == row['samples_sha256']
            else:
                mixture.ROOT = m.ROOT
                old_audit = audit.arm('raev2',stage,arm)
                # Reconcile actual Transformer and new-head calls as well as runtime counters.
                active_calls = c.read(m.ROOT/'preflight.json')['old_3k_replays'][0]['head_calls']['context3k']
                expected_heads = [0,0,0]
                if arm.startswith('context3k'):
                    expected_heads[0] = active_calls
                if arm.startswith('context20k'):
                    expected_heads[1] = active_calls
                for rec in row['records']:
                    with np.load(rec['file']) as batch:
                        np.testing.assert_array_equal(batch['block_calls'],np.full(30,100))
                        np.testing.assert_array_equal(batch['head_calls'],expected_heads)
            audits.append({**old_audit,'stage':stage,'arm':arm,'reused':reused})
            quality.append(dict(stage=stage,arm=arm,alpha=.39 if arm.endswith('half') else .78,
                fid=row['fid'],inception_score=row['inception_score'],samples=row['primary_samples'],
                reused=reused,seconds=row['seconds'],full_calls=row['full_calls_per_output'],
                prefix_calls=row['prefix_calls_at_inference'],samples_sha256=row['samples_sha256']))
    quality = pd.DataFrame(quality)
    decision = c.read(m.ROOT/'screen_decision.json')
    comparisons = []
    for suffix in ('base','half'):
        q = quality[quality.stage==m.SCREEN].set_index('arm')
        new = q.loc['context20k_'+suffix]
        native = q.loc['native_'+suffix]
        prior = q.loc['context_'+suffix]
        comparisons.append(dict(alpha=float(new.alpha),native_fid=float(native.fid),mlp_3k_fid=float(prior.fid),
            mlp_20k_fid=float(new.fid),delta_20k_vs_native=float(new.fid-native.fid),delta_20k_vs_3k=float(new.fid-prior.fid)))
    comparisons = pd.DataFrame(comparisons)
    frames = {'quality':quality,'paired_comparison':comparisons,'fixed_ema_loss':validation,
              'training_means_100':online_blocks,'recorded_training':history,'fixed_cases':pair_loss}
    for name,frame in frames.items():
        frame.to_csv(OUT/(name+'.csv'),index=False)
    with pd.ExcelWriter(OUT/'source_data.xlsx',engine='openpyxl') as writer:
        for name,frame in frames.items():
            frame.to_excel(writer,sheet_name=name,index=False)
            writer.book[name].freeze_panes='A2'
    book = load_workbook(OUT/'source_data.xlsx',data_only=True,read_only=True)
    for name,frame in frames.items():
        values = book[name].iter_rows(values_only=True)
        assert list(next(values)) == list(frame.columns)
        count = 0
        for expected,actual in zip(frame.itertuples(index=False,name=None),values):
            for x,y in zip(expected,actual):
                if pd.isna(x):
                    assert y is None
                elif isinstance(x,(int,float,np.number)) and not isinstance(x,(bool,np.bool_)):
                    assert np.isclose(x,y,rtol=1e-12,atol=1e-12)
                else:
                    assert x == y
            count += 1
        assert count == len(frame)
    book.close()
    plt.rcParams.update({'font.size':11,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none','pdf.fonttype':42})
    fig,axes = plt.subplots(1,2,figsize=(12.5,4.5),layout='constrained')
    axes[0].plot(online_blocks.last_step,online_blocks.context_mean,color='#21677F',lw=1.5)
    axes[0].axvline(3000,color='#777777',lw=1,ls=':')
    axes[0].set(title='Online training loss',xlabel='Optimizer step',ylabel='Clean-latent MSE',ylim=(0,1.05))
    axes[0].text(.98,.95,'Each point averages all 100 steps',transform=axes[0].transAxes,ha='right',va='top',fontsize=9)
    axes[1].plot(validation.step,validation.validation,color='#21677F',lw=2,marker='o',ms=3,label='Fixed validation: 1,000 cases')
    axes[1].plot(validation.step,validation.train,color='#B87932',lw=1.5,ls='--',marker='s',ms=3,label='Fixed training: 1,000 cases')
    axes[1].set(title='EMA on fixed clean / noise / time cases',xlabel='Optimizer step',ylabel='Clean-latent MSE',ylim=(0,.36))
    axes[1].legend(frameon=False,fontsize=9)
    for ax in axes:
        ax.grid(axis='y',color='#E5E8E9',lw=.7)
    for ext in ('png','svg','pdf'):
        fig.savefig(OUT/f'loss_curves.{ext}',dpi=170)
    plt.close(fig)
    for stage in stages:
        arms = (['native_base','context_base','context20k_base','native_half','context_half','context20k_half']
                if stage==m.SCREEN else c.read(m.ROOT/'raev2'/stage/'request.json')['arms'])
        labels = np.load(m.ROOT/'raev2'/stage/'inputs/labels.npy')[:4]
        fig,axes = plt.subplots(len(arms),4,figsize=(8.4,2.05*len(arms)),layout='constrained')
        for j,arm in enumerate(arms):
            with np.load(root_for(stage,arm)/'samples.npz') as d:
                imgs=d['arr_0'][:4]
            fid=float(quality[(quality.stage==stage)&(quality.arm==arm)].fid.iloc[0])
            suffix='0.39' if arm.endswith('half') else '0.78'
            name='Native' if arm.startswith('native') else ('20K MLP' if arm.startswith('context20k') else '3K MLP')
            for i in range(4):
                ax=axes[j,i]
                ax.imshow(imgs[i]);ax.set_xticks([]);ax.set_yticks([])
                for spine in ax.spines.values():spine.set_visible(False)
                if j==0:ax.set_title(f'Index {i}; class {labels[i]}',fontsize=9)
                if i==0:ax.set_ylabel(f'{name}\nalpha {suffix}\nFID {fid:.2f}',rotation=0,ha='right',va='center',fontsize=8)
        fig.suptitle('RAEv2: first four samples in fixed order',fontsize=12)
        fig.savefig(OUT/f'{stage}_first4.png',dpi=150)
        plt.close(fig)
    start=float(validation.iloc[0].validation)
    end=float(validation.iloc[-1].validation)
    change=end/start-1
    late=end/float(validation.loc[validation.step==15000,'validation'].iloc[0])-1
    best_native=float(comparisons.native_fid.min())
    best_20k=float(comparisons.mlp_20k_fid.min())
    display_validation=validation[validation.step.isin([3000,5000,10000,15000,20000])].rename(
        columns={'step':'训练步数','validation':'固定验证loss','train':'固定训练loss'})
    display_comparison=comparisons.rename(columns={'alpha':'额外系数alpha','native_fid':'原生IG FID',
        'mlp_3k_fid':'3K MLP FID','mlp_20k_fid':'20K MLP FID',
        'delta_20k_vs_native':'20K减原生','delta_20k_vs_3k':'20K减3K'})
    paragraphs=[
        '**RAEv2 Context MLP：延长至20K步的实际结果**',
        '已完成用户要求的总计20000步训练与两档IG系数生成复测。旧3000步EMA、31个原训练loss记录点与原256次抽样验证值全部精确复现，然后使用重建出的在线参数、AdamW状态和训练随机流继续17000步。3000步以后仅更新Context，未改变主干、读出结构、训练bank、损失、学习率或采样窗口。',
        f'在同一批400图上，20K原强度相对3K的FID变化为{comparisons.iloc[0].delta_20k_vs_3k:+.4f}，半强度为{comparisons.iloc[1].delta_20k_vs_3k:+.4f}。20K两档中最低FID为{best_20k:.4f}，原生IG两档中最低为{best_native:.4f}，差值为{best_20k-best_native:+.4f}。这些是该小规模配对筛选的观察值，不提供稳定收益或显著性结论。',
        f'新的固定1000例验证EMA MSE从3K时的{start:.6f}降至20K时的{end:.6f}，变化{change:+.2%}；最后15K到20K变化{late:+.2%}。原来的0.305192来自另一组256次随机抽取，不与新的固定1000例直接作跨设置差值。较低预测MSE不自动等于更好的生成质量。',
        '![训练与固定验证](data/raev2_context_20k_20260914/loss_curves.png)',
        display_validation.to_markdown(index=False,floatfmt=('.0f','.6f','.6f')),
        '图左每个点是完整100个训练步的平均，前3K来自此次精确重放；图右是每次相同clean/noise/time的1000个训练与1000个验证样本，各覆盖1000类。二者使用独立固定种子。验证没有更新参数，不干扰训练随机流；最终20K EMA在生成前已固定，没有根据FID选择中间checkpoint。',
        '**原400图配对复测，FID越低越好：**',
        display_comparison.to_markdown(index=False,floatfmt='.4f'),
        '这里只新生成20K两档共800张；原生IG和3K MLP两档1600张复用旧产物，其noise、标签及样本SHA已核对。另有8张旧3K前四样本的实现重放，latent与像素逐位一致，不作为新增质量样本。每图100次full、0额外prefix；各30个Transformer block实际调用100次，活动步MLP调用数另行核对。原IG scale=1.78等价于额外alpha=.78，半强度alpha=.39对应scale=1.39。',
        ('400图门槛未通过，未启动1K确认或新的系数搜索。' if not state['confirmed'] else
         '400图门槛通过后，按预先固定规则执行全新1000图确认；每类1图，四臂同噪声同标签。'),
        '400图仍只覆盖400类，门槛是本轮继续/停止规则，不是显著性检验。学习率保持3e-4，真实训练bank为每类5个、共5000个latent；即使训练到20K，也不等于已证明达到总体最优。']
    if state['confirmed']:
        paragraphs.extend(['**新1000图确认：**',quality[quality.stage==m.CONFIRM][['arm','fid','inception_score','samples']].to_markdown(index=False,floatfmt='.4f')])
    paragraphs.extend([
        '![固定400图前四个样本](data/raev2_context_20k_20260914/paired_400_first4.png)',
        f'训练记录总耗时{train["elapsed_seconds"]:.1f}秒，包含固定验证、特征缓存及checkpoint保存，不含此前准备过程。训练仅用GPU1，采样用GPU1–3，GPU0未使用。保留3K/5K/10K/15K/20K完整训练状态及最终20K EMA。',
        '[冻结协议](RAEV2_CONTEXT_20K_PROTOCOL_20260914_ZH.md) · [源数据工作簿](data/raev2_context_20k_20260914/source_data.xlsx) · [完整记录](data/raev2_context_20k_20260914/verification.json)。',
        '本轮回答延长训练是否改善这个具体弱读出及其IG效果；不将其作为共同误差分布或结构化mismatch的验证。长期核心方法目标未自动完成。'])
    REPORT.write_text('\n\n'.join(paragraphs)+'\n')
    process_checks=[]
    for stage in stages:
        workers=c.read(m.ROOT/f'{stage}_workers.json')
        for pid in [workers['parent']]+[r['pid'] for r in workers['children']]:
            p=Path('/proc')/str(pid)/'cmdline'
            command=p.read_bytes().decode(errors='replace') if p.exists() else ''
            assert 'experiments.raev2_context_20k_20260914.run' not in command
            assert 'experiments.raev2_context_20k_20260914.sample' not in command
            process_checks.append(dict(pid=pid,experiment_ended=True))
    verification=dict(passed=True,training_step=20000,replay_exact=True,source_workbook_reconciled=True,
        fixed_validation_cases=1000,validation_evaluations=18,new_quality_images=state['quality_images'],
        old_3k_replay_images=8,process_checks=process_checks,audits=audits,
        source_request_sha256=c.sha(tr.TRAIN/'request.json'),final_head_sha256=c.sha(tr.TRAIN/'head.pt'),
        visual_review_pending=not reviewed,goal_complete=False)
    c.atomic(OUT/'verification.json',verification)
    c.atomic(OUT/'manifest.json',dict(report_sha256=c.sha(REPORT),generator_sha256=c.sha(Path(__file__)),
        files={p.name:c.sha(p) for p in OUT.iterdir() if p.is_file() and p.name!='manifest.json'}))
    print('Report ready:',REPORT,'visual review pending:',not reviewed,flush=True)


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--visuals-reviewed',action='store_true')
    args=parser.parse_args()
    build(args.visuals_reviewed)
