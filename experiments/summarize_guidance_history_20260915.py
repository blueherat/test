"""Collect existing IG/readout/SG/weak-training results without launching jobs."""
import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

WORK = Path(__file__).resolve().parents[1]
DATA = WORK / 'docs/data'
EXPS = Path('/home/zhoushunyu/data/eqvae/experiments')
OUT = DATA / 'guidance_history_20260915'
REPORT = WORK / 'docs/GUIDANCE_COMPLETE_RESULTS_20260915_ZH.md'
ROWS, SOURCES, CHECKS = [], {}, []


def read(path):
    data = path.read_bytes()
    SOURCES[str(path)] = hashlib.sha256(data).hexdigest()
    return json.loads(data)


def number(value):
    return None if value in (None, '', 'None') else float(value)


def first(row, *keys, default=None):
    return next((row[k] for k in keys if row.get(k) not in (None, '', 'None')), default)


def add(raw, source, group, model, n=None, method=None, phase='', steps=None,
        coefficient=None, convention='', bank='', row_number=None):
    f = number(raw.get('fid'))
    valid = raw.get('valid', True) is not False and f is not None and math.isfinite(f)
    counts = raw.get('counts', {})
    n = first(raw, 'n', 'primary_samples', 'samples', 'sample_count', default=n)
    if isinstance(n, str) and not n.isdigit():
        n = None
    arm = method or first(raw, 'arm', 'method', 'branch', 'point', default='unknown')
    row = dict(group=group, model=model, bank=bank, phase=raw.get('stage', raw.get('phase', phase)),
        method=arm, samples=int(n) if n is not None else None,
        head_training_steps=first(raw, 'training_steps', default=steps),
        endpoint_step=raw.get('step'), weights=raw.get('weights'),
        coefficient_type=convention, coefficient=number(coefficient),
        alpha=number(raw.get('alpha')), omega=number(raw.get('omega')),
        seed=raw.get('seed'), kappa=number(raw.get('kappa')), time_shift=number(raw.get('shift')),
        solver=raw.get('solver'), sampling_steps=raw.get('steps'),
        fid=f, inception_score=number(raw.get('inception_score')), sfid=number(raw.get('sfid')),
        full_calls=first(raw, 'full_calls_per_output', 'full_calls', default=counts.get('full')),
        extra_prefix_calls=first(raw, 'prefix_calls_at_inference', 'prefix_calls', 'extra_prefix_calls'),
        seconds=number(raw.get('seconds')), valid=valid,
        first_1000_reused=raw.get('first_1000_reused'), reused=raw.get('reused'),
        checkpoint_sha256=raw.get('checkpoint_sha256'), samples_sha256=first(raw, 'samples_sha256', 'sample_sha256'),
        source_file=str(source), source_row=row_number, source_sha256=SOURCES[str(source)],
        source_output=first(raw, 'source_output', 'stored_at', 'samples_path', 'sample_path'))
    ROWS.append(row)


def csv_source(relative, group, model, n, bank):
    path = DATA / relative
    raw = path.read_bytes()
    SOURCES[str(path)] = hashlib.sha256(raw).hexdigest()
    for i, item in enumerate(csv.DictReader(raw.decode().splitlines()), 2):
        m = item.get('model', model)
        if m == 'sit': m = 'sit_small'
        arm = item.get('arm', '')
        coefficient, convention, steps = None, '', None
        if group.startswith('RAE'):
            coefficient = first(item, 'w', 'official_scale')
            if coefficient is None and item.get('alpha'): coefficient = float(item['alpha']) + 1
            convention = 'w: W+w*(S-W)'
            if 'context20k' in arm: steps = 20000
            elif 'context3k' in arm: steps = 3000
            elif arm.startswith(('mlp', 'native4')): steps = 50000
        elif group.startswith(('SiT_readout', 'SiT_matched', 'JiT_readout')):
            coefficient = .3 if m == 'jit' else .8
            if 'half' in arm: coefficient *= .5
            if 'double' in arm: coefficient *= 2
            if arm.startswith('strong'): coefficient = 0
            convention = 'a: S+a*(S-W)'
            steps = 3000 if arm in ('raw', 'mlp', 'context_base', 'native_fresh') else 50000
            if arm.startswith('cfg'):
                coefficient, convention, steps = 3, 'CFG scale', None
            if arm.startswith('strong'): steps = None
        elif group.startswith('IG_SG'):
            coefficient = 1 if ('sg_w1' in arm or 'log_k02_w1' in arm) else 0
            convention = 'omega: IG+omega*(S-Sref)'
        elif 'omega' in item:
            coefficient, convention = item.get('omega'), 'omega: base+omega*(S-Sref)'
        add(item, path, group, m, n, phase=item.get('stage', item.get('phase', '')),
            steps=steps, coefficient=coefficient, convention=convention,
            bank=bank, row_number=i)
        output = first(item, 'source_output', 'stored_at')
        if output:
            candidates = [Path(output) / 'metrics.json', Path(output) / 'fid.json']
            metric = next((p for p in candidates if p.exists()), None)
            if metric:
                check = read(metric)
                values = check if isinstance(check, list) else [check]
                diffs = [abs(float(v['fid']) - float(item['fid'])) for v in values if v.get('fid') is not None]
                assert diffs and min(diffs) < .001, (path, arm, diffs)
                CHECKS.append(dict(source=str(path), row=i, raw_metric=str(metric), fid_abs_error=min(diffs)))


def collect():
    specs = [
        ('context_reference_5k_20260912/selection_1k.csv','SiT_readout_1K','sit_small',1000,'sit_context_select1k'),
        ('context_reference_5k_20260912/quality.csv','SiT_readout_5K','sit_small',5000,'sit_context_independent5k'),
        ('ig_readout_matched_control_20260913/quality.csv','SiT_matched_3K','sit_small',1000,'sit_matched1k'),
        ('jit_readout_transfer_20260913/quality.csv','JiT_readout_1K','jit',1000,'jit_readout_select1k'),
        ('jit_readout_confirm_20260913/quality.csv','JiT_readout_5K','jit',5000,'jit_readout_independent5k'),
        ('raev2_context_20k_20260914/quality.csv','RAE_3K_20K_400','raev2',400,'rae_context400'),
        ('raev2_context_5k_20260914/quality.csv','RAE_context_5K','raev2',5000,'rae_heads5k'),
        ('raev2_shallow_ig_50k_20260914/results.csv','RAE_heads50K','raev2',None,'rae_heads_by_n'),
        ('raev2_shallow_ig_5k_selection_20260914/results.csv','RAE_heads5K_scan','raev2',5000,'rae_heads5k'),
        ('ig_sg_20260914/sit_small/results.csv','IG_SG_SiT','sit_small',1000,'ig_sg_seed2026091407'),
        ('ig_sg_20260914/jit/results.csv','IG_SG_JiT','jit',1000,'ig_sg_seed2026091407'),
        ('ig_sg_20260914/raev2/results.csv','IG_SG_RAE','raev2',1000,'ig_sg_seed2026091407'),
        ('ig_sg_20260914/confirm_1k/jit/results.csv','IG_SG_JiT_confirm','jit',1000,'ig_sg_seed2026091417'),
        ('self_guidance_20260913/screen_1k/results.csv','SG_initial_SiT','sit_small',1000,'sg_initial_screen1k'),
        ('self_guidance_20260913/confirm_1k/results.csv','SG_confirm_SiT','sit_small',1000,'sg_initial_confirm1k'),
        ('self_guidance_cross_model_20260913/screen_1k/jit/results.csv','SG_initial_JiT','jit',1000,'sg_cross_screen1k'),
        ('self_guidance_cross_model_20260913/screen_1k/raev2/results.csv','SG_initial_RAE','raev2',1000,'sg_cross_screen1k'),
        ('weak_reference_20260914/all_image_results.csv','SG_log_and_extensions','sit_small',1000,'see_phase_and_original_seed'),
    ]
    for args in specs: csv_source(*args)
    root = EXPS / 'guidance_loss_50k_20260914'
    for p in sorted(root.glob('sit_small/screen1000/*/metrics.json')):
        d = read(p); arm = p.parent.name
        alpha = .4 if arm == 'gaussian_half' else .8
        convention = 'a: S+a*(S-W), or corresponding residual coefficient'
        if arm == 'cfg_native': alpha, convention = 1.25, 'CFG extra coefficient'
        add(d, p, 'SiT_smallbank_fixed', 'sit_small', 1000,
            coefficient=alpha, convention=convention, bank='sit_common1k_2026091462')
    root = EXPS / 'guidance_strength_sweep_20260915'
    for p in sorted(root.glob('points/*/metrics.json')):
        d = read(p); arm, tick = d['point'].split('__c')
        add(d, p, 'SiT_smallbank_scan', 'sit_small', method=arm, steps=50000,
            coefficient=int(tick)/40, convention='a: S+a*(S-W)', bank='sit_common1k_2026091462')
    root = EXPS / 'guidance_dynamic_50k_20260915'
    for p in sorted(root.glob('*/points/*/n*/metrics.json')):
        d = read(p); arm, tick = d['point'].split('__c'); model = p.relative_to(root).parts[0]
        add(d, p, 'Full_data_50K', model, method=arm, steps=50000,
            coefficient=int(tick)/40, convention='a: S+a*(S-W)',
            bank='sit_common1k_2026091462' if d['n']==1000 and model=='sit_small' else f'{model}_dynamic5k')
    root = EXPS / 'adversarial_weak_training_20260915'
    for p in sorted(root.glob('endpoint_*/quality/*/c*/n*/metrics.json')):
        d = read(p); run = p.relative_to(root).parts[0]
        add(d, p, 'Endpoint_training', 'sit_small', method=run, steps=50000,
            coefficient=d['coefficient'], convention='a: S+a*(S-W)',
            bank='sit_common1k_2026091462' if d['n']==1000 else 'sit_small_dynamic5k')
    # Retain the unmodified CSV values but correct missing head-budget metadata
    # from the training protocols; strong/CFG control rows have no weak head.
    for row in ROWS:
        if row['method'] in ('strong', 'strong_e64'): row['head_training_steps'] = None
        if row['group']=='SiT_smallbank_fixed':
            if row['method']=='context': row['head_training_steps']=3000
            elif row['method']=='native': row['head_training_steps']=50000
        if row['bank']=='rae_heads_by_n':
            row['bank']='rae_heads5k' if row['samples']==5000 else 'rae_heads_balanced1k'
        if row['group']=='RAE_3K_20K_400' and row['method'] in ('context_base','context_half'):
            row['head_training_steps']=3000
        if row['group']=='Full_data_50K' and row['method']!='strong': row['weights']='ema'


def status():
    output=[]
    for p in sorted((EXPS/'adversarial_weak_training_20260915').glob('endpoint_*/progress.json')):
        d=read(p); run=p.parent.name; complete=p.parent/'complete.json'
        if complete.exists(): d.update(read(complete)); d['phase']='complete'
        log=p.parent/'train.jsonl'; updates=None
        if log.exists():
            entries=[json.loads(x) for x in log.read_text().splitlines() if x.strip()]
            updates=sum(bool(x.get('generator_updated')) for x in entries)
        output.append(dict(run=run,step=d.get('step'),phase=d.get('phase'),weak_updates_this_run=updates,
            target=d.get('target_step'),updated_utc=d.get('updated_utc'),source=str(p)))
    for p in sorted((EXPS/'guidance_dynamic_50k_20260915').glob('*/training/*/complete.json')):
        d=read(p); output.append(dict(run='full_data/'+d['model']+'/'+d['arm'],step=d['steps'],
            phase='complete',parameters=d.get('parameters'),source=str(p)))
    pause=read(EXPS/'ig_sg_5k_20260914/pause_snapshot.json')
    output.append(dict(run='IG_SG_5K',phase='paused',samples=pause['samples'],source=str(EXPS/'ig_sg_5k_20260914/pause_snapshot.json')))
    return output


def fmt(x):
    if x is None: return '—'
    if isinstance(x,float): return f'{x:.4f}'
    return str(x)


def table(rows, columns):
    lines=['| '+' | '.join(title for _,title in columns)+' |',
           '| '+' | '.join('---' for _ in columns)+' |']
    lines.extend('| '+' | '.join(fmt(row.get(k)).replace('|','/') for k,_ in columns)+' |' for row in rows)
    return '\n'.join(lines)


def write_outputs(snapshot, statuses):
    OUT.mkdir(parents=True, exist_ok=True)
    columns=list(ROWS[0])
    with (OUT/'all_results.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=columns);w.writeheader();w.writerows(ROWS)
    (OUT/'all_results.json').write_text(json.dumps(ROWS,ensure_ascii=False,indent=2)+'\n')
    (OUT/'status.json').write_text(json.dumps(statuses,ensure_ascii=False,indent=2)+'\n')
    wb=Workbook(); intro=wb.active;intro.title='说明'
    for line in [f'结果快照：{snapshot}','FID 越低越好；IS 越高越好。',
        '按 group/model/bank/samples 比较；不同样本数、输入组或模型不能直接混排。',
        'endpoint_step 是后训练全局步数，head_training_steps 是此前的基础头训练预算。',
        'RAEv2 的 w = 1+a；SG 表中的 omega 是额外系数，不是总权重 w。',
        '历史 CSV 可能重复引用同一批图；条目数不代表独立实验数或新生成图数。',
        'Full_data/Endpoint 的 5K 包含原 1K 选参样本；RAE 5K 后来用于选参。',
        'SiT/JiT 早期 MLP 的独立 5K 与筛选 1K 使用不同噪声。',
        '仅汇总既有结果；没有新增训练、生成或重新计算所有 FID。']:
        intro.append([line])
    intro.column_dimensions['A'].width=110
    display=['group','model','bank','phase','method','samples','head_training_steps','endpoint_step','weights',
             'coefficient_type','coefficient','alpha','omega','seed','kappa','time_shift','solver','sampling_steps','fid','inception_score','sfid','full_calls',
             'extra_prefix_calls','valid','first_1000_reused','reused','source_file','source_row','source_sha256','source_output']
    for title, rows in [('全部结果',ROWS)]+[(g,[r for r in ROWS if r['group']==g]) for g in dict.fromkeys(r['group'] for r in ROWS)]:
        ws=wb.create_sheet(title[:31]);ws.append(display)
        for row in rows: ws.append([row.get(k) for k in display])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='23445D')
        for i,k in enumerate(display,1):
            ws.column_dimensions[ws.cell(1,i).column_letter].width=24 if k not in ('source_file','source_output','coefficient_type') else 58
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value,float): cell.number_format='0.0000'
                cell.alignment=Alignment(vertical='top')
    wb.save(OUT/'results.xlsx')
    check=load_workbook(OUT/'results.xlsx',read_only=True,data_only=True)
    assert check['全部结果'].max_row==len(ROWS)+1
    check.close()

    lines=[f'结果快照：{snapshot}。本汇总覆盖本次讨论的 IG 读出头、冻结 MLP、SG/log-SG、弱分布构造与 endpoint 后训练实验。',
        f'共整理 {len(ROWS)} 条指标记录。含明确复用的历史结果，不代表 {len(ROWS)} 个独立实验。',
        '[完整 Excel](data/guidance_history_20260915/results.xlsx) · [完整 CSV](data/guidance_history_20260915/all_results.csv) · [来源与检查清单](data/guidance_history_20260915/manifest.json)',
        'FID 越低越好，IS 越高越好。不同模型、样本数量或噪声组分开报告；表中“最好”仅指已完成评价的点。',
        'RAEv2 的总外推系数 w 定义为 W+w(S−W)，不外推为 w=1；SiT/JiT 后训练扫描的额外系数 a 定义为 S+a(S−W)，不外推为 a=0。SG 的 omega 是叠加残差的额外系数，旧名字 w1 通常指 omega=1。',
        '早期 SiT/JiT MLP 的 5K 是独立噪声确认；RAEv2 头实验的同一 5K 后来用于选参，独立验证已取消；完整数据 SiT 与 endpoint 的 5K 包含原选参 1K。',
        '本次读取指标、扫描记录、训练完成标记和历史审计 CSV。对存在 source_output 的 CSV 逐项交叉核对原始指标；没有重新生成图像或全面重算特征。']
    titles={
        'SiT_readout_1K':'SiT：旧原生结构 IG 与 3K Context MLP，筛选 1K',
        'SiT_readout_5K':'SiT：旧 IG、Context MLP 与 ADG，独立 5K',
        'SiT_matched_3K':'SiT：相同 3K 训练预算的原结构头与 MLP，另一个 1K 输入组',
        'JiT_readout_1K':'JiT：原结构与 MLP，固定 1K',
        'JiT_readout_5K':'JiT：原结构与 MLP，独立 5K',
        'RAE_3K_20K_400':'RAEv2：早期 3K/20K 诊断，400 图',
        'RAE_context_5K':'RAEv2：原生头与 20K MLP，同组 5K',
        'RAE_heads50K':'RAEv2：50K 头的全部均衡 1K 与最初 5K（5K 后来归入选参）',
        'RAE_heads5K_scan':'RAEv2：所有已完成 5K 粗扫与第 8 层 MLP 细扫',
        'IG_SG_SiT':'SiT：IG 叠加 SG/log-SG，1K', 'IG_SG_JiT':'JiT：IG 叠加 SG/log-SG，1K',
        'IG_SG_RAE':'RAEv2：IG 叠加 SG/log-SG，1K','IG_SG_JiT_confirm':'JiT：IG+SG 独立第二种子 1K',
        'SG_initial_SiT':'SiT：最初原版 SG 与时间变体，1K','SG_confirm_SiT':'SiT：最初 SG 独立复核，1K',
        'SG_initial_JiT':'JiT：最初强模型/CFG 叠加 SG，1K','SG_initial_RAE':'RAEv2：最初强模型/IG 叠加 SG，1K',
        'SG_log_and_extensions':'log-score、带宽、角度与矩匹配等历史 1K 实验，按 model/phase 分开比较',
        'SiT_smallbank_fixed':'SiT：固定小数据集的 50K 弱头，固定系数 1K',
        'SiT_smallbank_scan':'SiT：固定小数据集的外推系数扫描，1K',
        'Full_data_50K':'完整数据动态采样的 50K 重训与全部系数扫描',
        'Endpoint_training':'GAN / moments / joint moments / energy 后训练的全部已落盘评价',
    }
    for group in dict.fromkeys(r['group'] for r in ROWS):
        rows=[r for r in ROWS if r['group']==group]
        lines.append('**'+titles[group]+'**')
        cols=[('method','方法'),('samples','采样数')]
        if group=='SG_log_and_extensions': cols=[('model','模型'),('phase','阶段')]+cols
        if group=='Endpoint_training':cols += [('endpoint_step','全局步'),('weights','权重')]
        cols += [('coefficient','系数'),('fid','FID'),('inception_score','IS')]
        if group.startswith(('SiT_readout','JiT_readout','IG_SG','SG_')):cols += [('full_calls','主干调用/图')]
        lines.append(table(rows,cols))
    lines += ['**训练和暂停状态**',table(statuses,[('run','实验'),('step','步数'),('phase','状态'),('weak_updates_this_run','本段弱头更新次数')]),
        'IG+SG 的 5K 队列实际仅保留 RAEv2 log-SG、w_SG=1.2 的 2,272 张，未完成 5K，无有效 5K FID；不把计划文档里的目标数量当作完成结果。',
        'RAEv2 第 4 层原结构头 w=3 的均衡 1K 采样发生数值失败，未提供部分 FID。第 4 层后续搜索、第 8 层 MLP 的 w=2 及额外细扫点按用户指示取消，缺失位置不填估计值。原生头和第 4 层原结构头已完成的 w=2 结果仍然保留。',
        'SiT 的小 bank real50K 只使用 1,800 张固定 clean latent；旧 Context3K 使用 25,600 张图的动态 posterior；完整数据重训恢复全部训练数据且改变优化设置。因此这些并非只改变 loss 或训练时长的结构消融。',
        'moments 第 600 步在线头在 1K 小胜但 5K 落后；joint moments 第 1,000 步的最好已测 5K 为 36.8117，仍未超过同组 guided_weak 的 36.6888。energy 的数值以本快照已完成指标为准。',
        '同样的 5K 规模不代表同一噪声组；例如旧 SiT Context 的 37.2343 与新 real50K 的 36.9095，不能当作配对训练改进量。',
        '[SiT 头的独立 5K 原报告](CONTEXT_REFERENCE_5K_RESULTS_20260912_ZH.md) · [JiT 头的独立 5K](JIT_READOUT_CONFIRM_RESULTS_20260913_ZH.md) · [RAEv2 5K 扫描原报告](RAEV2_SHALLOW_IG_5K_SELECTION_RESULTS_20260914_ZH.md) · [1K/5K 反转审计](ENDPOINT_MOMENTS_1K_5K_AUDIT_20260915_ZH.md)']
    REPORT.write_text('\n\n'.join(lines)+'\n')
    manifest=dict(snapshot=snapshot,rows=len(ROWS),groups=dict(Counter(r['group'] for r in ROWS)),
        source_files=SOURCES,csv_raw_cross_checks=CHECKS,
        validation=dict(all_finite_valid_fids=all(r['fid'] is not None and math.isfinite(r['fid']) for r in ROWS if r['valid']),
                        workbook_roundtrip_rows=len(ROWS),regenerated_images=0),
        outputs={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in (REPORT,OUT/'all_results.csv',OUT/'all_results.json',OUT/'results.xlsx',OUT/'status.json')})
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(snapshot=snapshot,rows=len(ROWS),groups=manifest['groups'],
                         source_files=len(SOURCES),cross_checks=len(CHECKS),report=str(REPORT)),ensure_ascii=False))


if __name__=='__main__':
    collect()
    snapshot=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    write_outputs(snapshot,status())
