"""Build a source-backed report after the fixed generation comparison finishes."""
import argparse
from pathlib import Path
import json
import re
import textwrap
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from . import core as m

c=m.c
OUT=c.WORK/'docs/data/transition_decontamination_20260913'
REPORT=c.WORK/'docs/TRANSITION_DECONTAMINATION_RESULTS_20260913_ZH.md'
LABELS={'ode':'Original ODE','ode_competitor':'ADG / APG ODE','sde':'Affine guided SDE',
        'positive':'Positive density difference','mean_gaussian':'Same mean, isotropic Gaussian',
        'moment_gaussian':'Same mean and covariance Gaussian'}


def build(reviewed=False):
    m.configure()
    status=c.read(m.ROOT/'status.json')
    assert status['phase']=='complete' and status['images']==4800
    audit=c.read(m.ROOT/'audit.json');assert audit['passed'] and audit['samples']==4800
    decision=c.read(m.ROOT/'decision.json')
    timing=c.read(m.ROOT/'benchmark.json');assert timing['complete']
    numeric=c.read(m.ROOT/'numerical_checks.json');assert numeric['passed']
    assert numeric['kernel_source_sha256']==c.sha(m.kernel.__file__)
    preflight=c.read(m.ROOT/'implementation_checks.json');assert preflight['passed']
    OUT.mkdir(parents=True,exist_ok=True)
    rows,costs,sources=[],[],[]
    for track in m.TRACKS:
        rp,request=m.verify(track)
        root=m.ROOT/'sit_small'/m.stage(track)
        values=c.read(root/'results.json')
        base=next(r['fid'] for r in values if r['arm']=='ode')
        for value in values:
            name=value['arm'];out=root/name
            assert c.sha(out/'samples.npz')==value['samples_sha256']
            for record in value['records']:
                assert c.sha(record['file'])==record['sha256']
            median=timing['tracks'][track]['medians'][name]
            row=dict(track=track,arm=name,fid=value['fid'],inception_score=value['inception_score'],
                delta_from_ode=value['fid']-base,samples=value['primary_samples'],
                full_calls=value['full_calls_per_output'],prefix_calls=value['prefix_calls_at_inference'],
                median_seconds=median)
            rows.append(row)
            for repeat,seconds in enumerate(timing['tracks'][track]['seconds'][name],1):
                costs.append(dict(track=track,arm=name,repeat=repeat,seconds=seconds,batch=timing['batch']))
            for p in (out/'metrics.json',out/'summary.json',out/'samples.npz',out/'activations.npz'):
                sources.append(dict(title=f'{track} / {name} / {p.name}',path=str(p),sha256=c.sha(p)))
        sources.append(dict(title=track+' frozen request',path=str(rp),sha256=c.sha(rp)))
    for p in (m.PROTOCOL,m.ROOT/'audit.json',m.ROOT/'benchmark.json',m.ROOT/'numerical_checks.json',
              m.ROOT/'kernel_lut.npz',m.ROOT/'implementation_checks.json'):
        sources.append(dict(title=p.name,path=str(p),sha256=c.sha(p)))
    papers=c.read(c.WORK/'readings/transition_decontamination_20260913/source_manifest.json')
    for item in papers:
        assert c.sha(c.WORK/item['path'])==item['sha256']
        sources.append(dict(title=item['title'],path=item['url'],sha256=item['sha256']))
    quality=pd.DataFrame(rows)
    costs=pd.DataFrame(costs)
    theory=pd.DataFrame(numeric['quadrature_cases'])
    decisions=pd.DataFrame([dict(track=k,passes_gate=v['passes_gate'],candidate_fid=v['candidate_fid'],
        best_control=v['best_control'],best_control_fid=v['best_control_fid'],difference=v['difference']) for k,v in decision.items()])
    frames={'quality':quality,'timing':costs,'kernel_moments':theory,'decisions':decisions,'sources':pd.DataFrame(sources)}
    for name,df in frames.items():df.to_csv(OUT/(name+'.csv'),index=False)
    workbook=OUT/'source_data.xlsx'
    with pd.ExcelWriter(workbook,engine='openpyxl') as writer:
        for name,df in frames.items():df.to_excel(writer,sheet_name=name,index=False)
        for ws in writer.book:
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.fill=PatternFill('solid',fgColor='252525');cell.font=Font(color='FFFFFF',bold=True)
            for col in ws.columns:
                length=max(len(str(x.value or '')) for x in col)
                ws.column_dimensions[col[0].column_letter].width=min(70,max(14,length+2))
    check=load_workbook(workbook,data_only=True)
    for name,df in frames.items():
        values=list(check[name].values)
        assert list(values[0])==list(df.columns) and len(values)==len(df)+1
        for expected,actual in zip(df.itertuples(index=False,name=None),values[1:]):
            for x,y in zip(expected,actual):
                if isinstance(x,(int,float,np.number)) and not isinstance(x,(bool,np.bool_)):
                    assert np.isclose(x,y,rtol=1e-12,atol=1e-12)
                else:assert x==y,(name,x,y)
    plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False})
    fig,axes=plt.subplots(1,2,figsize=(13.5,4.8),layout='constrained')
    for ax,track in zip(axes,m.TRACKS):
        data=quality[quality.track==track]
        ax.barh(np.arange(len(data)),data.delta_from_ode,color=['#21677F' if a=='positive' else '#999999' for a in data.arm])
        ax.set_yticks(np.arange(len(data)),[LABELS[a].replace('ADG / APG','ADG' if track=='ig' else 'APG') for a in data.arm])
        ax.invert_yaxis();ax.axvline(0,color='black',lw=.8)
        span=max(float(data.delta_from_ode.max()-data.delta_from_ode.min()),1.)
        left=min(0.,float(data.delta_from_ode.min()))-.1*span
        right=max(0.,float(data.delta_from_ode.max()))+.65*span
        ax.set_xlim(left,right)
        for j,row in enumerate(data.itertuples()):
            ax.text(max(0.,row.delta_from_ode)+.025*span,j,f'FID {row.fid:.2f}',va='center',fontsize=9)
        ax.set_title(track.upper()+': 400 paired images')
        ax.set_xlabel('FID difference from original ODE (lower is better)')
    for suffix in ('png','svg','pdf'):fig.savefig(OUT/f'quality_comparison.{suffix}',dpi=170)
    plt.close(fig)
    for track in m.TRACKS:
        root=m.ROOT/'sit_small'/m.stage(track)
        labels=np.load(root/'inputs/labels.npy')[:4]
        fig,axes=plt.subplots(6,4,figsize=(8.4,12.2),layout='constrained')
        for j,arm in enumerate(m.ARMS):
            with np.load(root/arm/'samples.npz') as d:images=d['arr_0'][:4]
            fid=float(quality[(quality.track==track)&(quality.arm==arm)].fid.iloc[0])
            for i in range(4):
                axes[j,i].imshow(images[i]);axes[j,i].set_xticks([]);axes[j,i].set_yticks([])
                for spine in axes[j,i].spines.values():spine.set_visible(False)
                if j==0:axes[j,i].set_title(f'Index {i}; class {labels[i]}',fontsize=9)
                if i==0:
                    label=LABELS[arm].replace('ADG / APG','ADG' if track=='ig' else 'APG')
                    axes[j,i].set_ylabel(textwrap.fill(label,width=22)+f'\nFID {fid:.2f}',rotation=0,ha='right',va='center',fontsize=8)
        fig.suptitle(track.upper()+': first four samples in order; no image selection',fontsize=11)
        fig.savefig(OUT/f'{track}_first4.png',dpi=150);plt.close(fig)
    passed=[t for t,d in decision.items() if d['passes_gate']]
    lead=('两条路线均未通过固定筛选，停止这一有限转移构造。' if not passed else
          '固定400图筛选通过的路线为'+', '.join(passed)+'；尚需独立1K确认。')
    paragraphs=['**有限随机转移的正性去污染：CFG 与 IG 固定生成实验**',lead,
        '本轮完成12组、每组400张，共4800张全新图像。IG使用已验证的MLP弱读出，CFG使用原条件及无条件预测。没有新增训练、弱主干、来源分类器或额外前缀，也没有根据生成结果修改参数。两个噪声bank独立，分别使用seed2026121431与2026121433，均衡100个本地类别。',
        '![同预算比较](data/transition_decontamination_20260913/quality_comparison.png)',
        quality[['track','arm','fid','inception_score','samples','full_calls','prefix_calls']].to_markdown(index=False,floatfmt='.4f'),
        '每条路线内部所有随机臂的初始噪声、类别与128步Gaussian随机数配对。原ODE及竞争基线使用64步Heun；随机臂使用128步Euler。IG各臂128次full，CFG各臂224次full，额外prefix均为0。ODE与SDE不同，因此收益必须同时超过普通guided SDE，不能只与ODE比较。CFG竞争基线是仓库既有APG velocity历史适配，IG竞争基线是同MLP读出的ADG。',
        decisions.to_markdown(index=False,floatfmt='.4f'),
        '门槛要求候选比同路线全部五个控制至少低2 FID，且IS达到原ODE的90%。这是固定继续/停止规则，不是统计显著性检验。400图只提供筛选证据，不能和不同样本数或旧bank的绝对FID交叉比较。']
    for track in m.TRACKS:
        d=decision[track]
        paragraphs.append(f"{track.upper()}候选相对原ODE的FID差为{d['differences']['ode']:+.4f}，相对普通SDE为{d['differences']['sde']:+.4f}，相对均值匹配Gaussian为{d['differences']['mean_gaussian']:+.4f}，相对均值及协方差匹配Gaussian为{d['differences']['moment_gaussian']:+.4f}。这些对照用于区分有限转移分布形状、平均漂移与二阶矩；不将候选和矩对照接近自动解释成某一个机制已被识别。")
    paragraphs += [
        '候选把原来的密度去污染动机放到每一步的条件转移上。对于两个近似Gaussian转移S、W，定义\n\n\\[Q=\\frac{[S-\\kappa W]_+}{\\int[S-\\kappa W]_+}.\\]\n\n它是有符号反演A=(S-kappa W)/(1-kappa)的一个最近TV概率修复：A的负质量给出任何概率修复所需TV改变量的下界，归一化正部分达到此下界，但解不必唯一。该性质保证合法转移，不保证接近真实数据。',
        '这是一项额外的转移核假设，不能由端点共同污染自动推出。不同均值、相同协方差的两个Gaussian本来就不能满足全空间正比例去污染；正部分操作承认并修复这一失配。其修复范围限于正性，没有识别真实错误成分、卷积核或一般加性误差。',
        '采样时令mu_i=z+h*((1+t*(1-t))*v_i-(1-t)*z)，sigma=sqrt(2*h)*(1-t)。这是选择扩散方差率2*(1-t)^2后对相应SDE做Euler离散。在精确共同Gaussian前向族条件下，所加漂移与扩散保持同一边缘；冻结近似网络和有限步离散不保证该等价。[1]',
        '因为两个转移协方差相同，密度比只依赖均值差方向上的一个标量。令d=||mu_w-mu_s||/sigma，kappa=alpha/(1+alpha)，则标准化标量密度正比于[phi(u)-kappa*phi(u-d)]_+，支持u<c=d/2-log(kappa)/d。其余方向保留标准Gaussian。它改变完整转移分布；均值匹配与协方差匹配两臂是独立的控制。',
        '固定维数且步长趋零时，d=O(sqrt(h))，正部分截断消失，均值回到普通affine guidance，协方差改变量为O(h^2)。因此连续极限仍为普通guided SDE；本次研究的是有限步修复，而非新的连续极限理论。原alpha和窗口仅用于小步幅度匹配，不能称为估计到污染比例。',
        f"逆CDF采用三个固定1025×1025 FP32表，d轴按sqrt(d)取网格。独立数值积分核验了21个归一化、均值、方差及TV距离案例，并用100000次随机抽样检查采样分布。18000个抽查点的最大标准化偏移误差为{max(numeric['lookup_max_abs_shift_errors']):.6g}。标准Gaussian投影超出±7的概率约{numeric['gaussian_projection_tail_bound']:.3g}；表范围外采用截断数值约定，不能称无限精度精确采样。",
        '原IG、CFG的完整ODE轨迹与旧实现逐项相同，CFG的APG轨迹也逐项相同。所有随机臂的零引导分支逐项相同，普通SDE与独立显式漂移重放逐项相同。正式每批保存全部层调用数、头调用数、噪声与类别、随机增量种子、FP32终态和像素；权重、源代码及输入请求hash核验通过。',
        '同卡轮换预热后各计时三次，batch8，包含采样、VAE解码及像素转换。下列时间与并行质量采样日志分开；只有三次重复，不给精确延迟置信区间。',
        quality[['track','arm','median_seconds']].to_markdown(index=False,floatfmt='.6f'),
        f"候选没有新增可训练参数，查找表占{timing['lookup_bytes']/2**20:.2f} MiB。IG与自己的MLP基线相比没有新增读出参数；CFG算法不调用弱头。比较运行时为共用两条路线加载了MLP与全部查找表，其总显存不能冒充纯CFG最小部署显存。",
        f"全部4800张及批次元数据已核对。对同一ADM缓存特征使用FP64对称Gram公式复算FID，最大绝对差{audit['max_fid_error']:.6g}。这是同特征复算，不是另一套独立特征提取器验证。",
        '![IG固定前四张](data/transition_decontamination_20260913/ig_first4.png)',
        '![CFG固定前四张](data/transition_decontamination_20260913/cfg_first4.png)',
        '线性污染和自适应系数已有Feedback Guidance先例；概率流叠加和相应SDE族已有SuperDiff等先例。[1][2] 本次结果不建立新颖性，也不把有限转移密度的合法性等同于生成质量。完整误差假设仍需同时处理实际分布错配与网络近似误差。',
        '[冻结协议](TRANSITION_DECONTAMINATION_PROTOCOL_20260913_ZH.md) · [采样实现](../experiments/transition_decontamination_20260913/core.py) · [转移核实现](../experiments/transition_decontamination_20260913/kernel.py) · [源数据工作簿](data/transition_decontamination_20260913/source_data.xlsx) · [核验记录](data/transition_decontamination_20260913/verification.json)。',
        '[1] Skreta 等，The Superposition of Diffusion Models Using the Itô Density Estimator，2025-02-28，Proposition 1及§2–3：[原文](https://arxiv.org/html/2412.17762v2)。',
        '[2] Koulischer 等，Feedback Guidance of Diffusion Models，2025，§3：[原文](https://arxiv.org/html/2506.06085v2)。'
    ]
    REPORT.write_text('\n\n'.join(paragraphs)+'\n')
    # This check covers generated local links; web claims are traced to the archived primary papers.
    for link in re.findall(r'\]\(([^)]+)\)',REPORT.read_text()):
        if not link.startswith('http') and not link.endswith('verification.json'):
            assert (REPORT.parent/link).exists(),link
    source_files={str(p):c.sha(p) for p in (m.ROOT/'audit.json',m.ROOT/'decision.json',m.ROOT/'benchmark.json',
        m.ROOT/'numerical_checks.json',m.ROOT/'implementation_checks.json',Path(m.__file__),Path(m.kernel.__file__),m.PROTOCOL)}
    verification=dict(passed=True,visual_inspection_pending=not reviewed,samples=4800,arms=12,
        workbook_tables_reconciled=True,all_requests_verified=True,source_files=source_files,
        max_fid_error=audit['max_fid_error'],goal_complete=False)
    c.atomic(OUT/'verification.json',verification)
    manifest=dict(report_sha256=c.sha(REPORT),generator_sha256=c.sha(Path(__file__)),
        files={str(p.relative_to(c.WORK)):c.sha(p) for p in sorted(OUT.iterdir()) if p.name!='artifact_manifest.json'})
    c.atomic(OUT/'artifact_manifest.json',manifest)
    print('Final report verification passed' if reviewed else 'Built report; visual inspection pending',REPORT,flush=True)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--visuals-reviewed',action='store_true');a=p.parse_args()
    build(a.visuals_reviewed)
