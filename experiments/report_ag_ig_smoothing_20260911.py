#!/usr/bin/env python3
"""Exact controls and inspectable plots for the AG/IG smoothing diagnostic."""
from __future__ import annotations
import hashlib
import json
from pathlib import Path
import sys

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy.optimize import minimize_scalar
from scipy.special import logsumexp

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/data/ag_ig_smoothing_20260911'


def gmm(x, u, base=(.08, .30), means=(1., .6)):
    x = np.asarray(x, dtype=np.float64)
    m = np.asarray(means)
    v = np.asarray(base) + u
    a = np.sum(x*m/v, axis=-1, keepdims=True)
    s = -x/v + m/v*np.tanh(a)
    du = x/v**2 - m/v**2*np.tanh(a) - m/v*(1-np.tanh(a)**2)*np.sum(x*m/v**2, axis=-1, keepdims=True)
    component = -.5*np.sum((x[..., None, :] - np.stack([-m, m]))**2/v, axis=-1)
    component -= .5*np.sum(np.log(2*np.pi*v))
    logp = logsumexp(component, axis=-1)-np.log(2.)
    return logp, s, du


def exact_controls():
    OUT.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(20260911)
    q = .10
    signs = rng.choice([-1., 1.], size=(4096, 1))
    points = signs*np.array([1., .6]) + rng.normal(size=(4096, 2))*np.sqrt(np.array([.08,.30])+q)
    _, s, ds = gmm(points, q)
    rows = []
    fit_err = []
    for tau in [.002,.01,.03,.1,.3,1.,2.]:
        _, sw, _ = gmm(points, q+tau)
        gap = sw-s
        objective = lambda a: float(np.square(gmm(points, q+a)[1]-sw).mean())
        fitted = minimize_scalar(objective, bounds=(0.,3.), method='bounded', options={'xatol':1e-12})
        cosine = float(np.sum(gap*ds)/np.sqrt(np.sum(gap**2)*np.sum(ds**2)))
        point_cos = np.sum(gap*ds,axis=1)/np.sqrt(np.sum(gap**2,axis=1)*np.sum(ds**2,axis=1))
        residual = objective(fitted.x)/float(np.square(gap).mean())
        fit_err.append(abs(fitted.x-tau))
        rows.append(dict(tau=tau, fitted_tau=fitted.x, relative_fit_residual=residual,
                         derivative_cosine=cosine, pointwise_cosine_median=float(np.median(point_cos)),
                         pointwise_cosine_q10=float(np.quantile(point_cos,.1))))
    pd.DataFrame(rows).to_csv(OUT/'toy_finite_smoothing.csv', index=False)
    # Coordinate transformation independently reconstructed from an exact denoiser.
    t = 1/(1+np.sqrt(q))
    z = points*t
    clean = points+q*s
    velocity = (clean-z)/(1-t)
    transformed = t*(t*velocity-z)/(1-t)
    coordinate_error = float(np.max(np.abs(transformed-s)))
    h=1e-5
    numeric_du = (gmm(points,q+h)[1]-gmm(points,q-h)[1])/(2*h)
    derivative_error = float(np.sqrt(np.mean((numeric_du-ds)**2)/np.mean(ds**2)))
    # Frozen-noise density illustration. This is NOT an endpoint sampling law.
    x = np.linspace(-4.,4.,4001)
    q0, extra_strong, tau0, gamma = .02,.08,1.,1.
    lp_star = gmm(x[:,None],q0,base=(.0256,),means=(1.,))[0]
    lp_s = gmm(x[:,None],q0+extra_strong,base=(.0256,),means=(1.,))[0]
    lp_w = gmm(x[:,None],q0+extra_strong+tau0,base=(.0256,),means=(1.,))[0]
    log_power = (1+gamma)*lp_s-gamma*lp_w
    unnormalized = np.exp(log_power-log_power.max())
    p_power = unnormalized/np.trapezoid(unnormalized,x)
    score_s = gmm(x[:,None],q0+extra_strong,base=(.0256,),means=(1.,))[1][:,0]
    score_w = gmm(x[:,None],q0+extra_strong+tau0,base=(.0256,),means=(1.,))[1][:,0]
    component_variance=.0256+q0+extra_strong
    total_variance=1.+component_variance
    log_coarse=-.5*x*x/total_variance-.5*np.log(2*np.pi*total_variance)
    log_coarse_power=(1+gamma)*lp_s-gamma*log_coarse
    coarse_power=np.exp(log_coarse_power-log_coarse_power.max())
    coarse_power/=np.trapezoid(coarse_power,x)
    density = pd.DataFrame(dict(x=x, target=np.exp(lp_star), strong=np.exp(lp_s), weak=np.exp(lp_w),
                                guided_power=p_power, density_ratio=np.exp(lp_s-lp_w), guidance=score_s-score_w,
                                moment_matched_weak=np.exp(log_coarse),moment_matched_guided_power=coarse_power))
    density.to_csv(OUT/'toy_density.csv',index=False)
    moments=[dict(distribution='strong_two_modes',mean=0.,variance=total_variance,
                  fourth_cumulant=-2.,component_variance=component_variance),
             dict(distribution='gaussian_heat_weak',mean=0.,variance=total_variance+tau0,
                  fourth_cumulant=-2.,component_variance=component_variance+tau0),
             dict(distribution='moment_matched_coarse_weak',mean=0.,variance=total_variance,
                  fourth_cumulant=0.,component_variance=total_variance)]
    pd.DataFrame(moments).to_csv(OUT/'toy_coarsening_moments.csv',index=False)
    i0 = len(x)//2
    ipos = int(np.argmin(abs(x-1.)))
    mode = float(x[x>0][np.argmax(p_power[x>0])])
    # Exact zero-error AG coefficients for a Gaussian target of variance 1.
    # eps_s=q, eps_w=2q: eps_s -> 0, yet required gamma -> 1, not 0.
    limits=[]
    for qn in np.geomspace(1e-6,10.,141):
        gs = qn; gw=2*qn
        g = gs*(1+qn+gw)/((gw-gs)*(1+qn))
        s_star = -1/(1+qn)
        s_s=-1/(1+qn+gs); s_w=-1/(1+qn+gw)
        err=abs(s_s+g*(s_s-s_w)-s_star)
        limits.append(dict(q=qn,eps_strong=gs,eps_weak=gw,gamma_exact=g,
                           correction_score_at_y1=g*(s_s-s_w),score_identity_error=err))
    pd.DataFrame(limits).to_csv(OUT/'toy_low_noise_limit.csv',index=False)
    audit=dict(coordinate_identity_max_error=coordinate_error,
               derivative_relative_rms_error=derivative_error,
               finite_tau_recovery_max_abs_error=max(fit_err),
               finite_fit_max_relative_residual=max(r['relative_fit_residual'] for r in rows),
               strong_valley_to_at_x1=float(np.exp(lp_s[i0]-lp_s[ipos])),
               weak_valley_to_at_x1=float(np.exp(lp_w[i0]-lp_w[ipos])),
               power_valley_to_at_x1=float(p_power[i0]/p_power[ipos]),
               positive_guided_power_mode=mode,
               moment_matched_valley_to_at_x1=float(np.exp(log_coarse[i0]-log_coarse[ipos])),
               moment_matched_guided_valley_to_at_x1=float(coarse_power[i0]/coarse_power[ipos]),
               moment_matched_variance_difference=0.,
               gamma_at_smallest_q=limits[0]['gamma_exact'],
               gaussian_score_identity_max_error=max(r['score_identity_error'] for r in limits))
    assert coordinate_error < 1e-12
    assert derivative_error < 1e-7
    assert max(fit_err)<1e-7
    assert audit['finite_fit_max_relative_residual']<1e-10
    assert audit['gaussian_score_identity_max_error']<1e-12
    (OUT/'toy_audit.json').write_text(json.dumps(audit,indent=2)+'\n')
    return audit


def figures():
    plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,
                         'figure.dpi':150,'savefig.dpi':180})
    density=pd.read_csv(OUT/'toy_density.csv')
    finite=pd.read_csv(OUT/'toy_finite_smoothing.csv')
    limit=pd.read_csv(OUT/'toy_low_noise_limit.csv')
    fig, axes=plt.subplots(1,3,figsize=(14,4.9),layout='constrained')
    ax=axes[0]
    for name,color,label in [('strong','#146b8f','Strong'),('weak','#da9549','Weak = heat(strong)'),('guided_power','#943b72','Fixed-noise guided density')]:
        ax.plot(density.x,density[name],color=color,label=label)
    ax.set(xlim=(-2.5,2.5),xlabel='x',ylabel='Density',title='A. Heat smoothing can join two modes')
    ax.legend(fontsize=8,loc='upper center',bbox_to_anchor=(.5,-.19))
    ax=axes[1]
    ax.semilogx(finite.tau,finite.derivative_cosine,'o-',color='#146b8f')
    ax.set(xlabel='Known extra variance (tau)',ylabel='Cosine with noise derivative',ylim=(.65,1.01),
           title='B. Exact heat family need not be collinear')
    ax.text(.04,.13,'Finite-scale fit residual < 1e-10\nfor every point',transform=ax.transAxes,fontsize=9)
    ax=axes[2]
    ax.semilogx(limit.q,limit.gamma_exact,color='#943b72')
    ax.set(xlabel='Noise variance q (decreases in sampling)',ylabel='Exact score correction coefficient',
           title='C. Low-noise coefficient need not vanish',ylim=(0,3.05))
    ax.axhline(1,color='gray',ls='--',lw=.8)
    ax.text(.05,.08,'eps(strong)=q; eps(weak)=2q\ngamma -> 1 although both errors vanish',transform=ax.transAxes,fontsize=8)
    fig.savefig(OUT/'exact_smoothing_controls.png');fig.savefig(OUT/'exact_smoothing_controls.pdf');plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(10.8,4.9),layout='constrained')
    for ax in axes:ax.set(xlim=(-2.8,2.8),xlabel='x',ylabel='Density')
    for name,color,style,label in [('strong','#146b8f','-','Strong: two modes'),
                                  ('weak','#da9549','--','Gaussian heat weak'),
                                  ('moment_matched_weak','#48986a','-.','Moment-matched single-mode weak')]:
        axes[0].plot(density.x,density[name],color=color,ls=style,label=label)
    axes[0].set_title('A. Two different ways to merge modes')
    axes[0].legend(fontsize=8,loc='upper center',bbox_to_anchor=(.5,-.19))
    for name,color,label in [('strong','#146b8f','Strong'),('guided_power','#da9549','Guidance against heat weak'),
                            ('moment_matched_guided_power','#48986a','Guidance against moment-matched weak')]:
        axes[1].plot(density.x,density[name],color=color,label=label)
    axes[1].set_title('B. Both density ratios can sharpen modes')
    axes[1].legend(fontsize=8,loc='upper center',bbox_to_anchor=(.5,-.19))
    fig.savefig(OUT/'coarsening_without_convolution.png');fig.savefig(OUT/'coarsening_without_convolution.pdf');plt.close(fig)
    if (OUT/'neural_summary.csv').exists():
        table=pd.read_csv(OUT/'neural_summary.csv')
        fig,axes=plt.subplots(1,2,figsize=(12,4.3),layout='constrained')
        names=['depth4','depth6','depth8','depth10','external_v500']
        ts=sorted(table.t.unique())
        for ax,kind,title in zip(axes,['finite_nonnegative','derivative_nonnegative'],
                                 ['A. Best nonnegative finite heat shift','B. Best nonnegative noise derivative']):
            part=table[table.fit_type==kind].pivot(index='weak',columns='t',values='test_ratio').loc[names,ts]
            im=ax.imshow(part.to_numpy(),vmin=0,vmax=1.05,cmap='YlOrRd',aspect='auto')
            ax.set_xticks(range(len(ts)),[f'{t:g}' for t in ts]);ax.set_yticks(range(len(names)),names)
            ax.set(xlabel='Flow time t (noise -> data)',title=title)
            for i in range(len(names)):
                for j in range(len(ts)):
                    value=part.iloc[i,j]
                    ax.text(j,i,f'{value:.3f}',ha='center',va='center',color='white' if value>.65 else 'black')
        fig.colorbar(im,ax=axes,label='Held-out residual / original strong-weak gap energy',shrink=.85)
        fig.savefig(OUT/'neural_heat_shift_residuals.png');fig.savefig(OUT/'neural_heat_shift_residuals.pdf');plt.close(fig)
        if (OUT/'refined_summary.csv').exists():
            refined=pd.read_csv(OUT/'refined_summary.csv')
            fig,axes=plt.subplots(1,2,figsize=(12,4.3),layout='constrained')
            for ax,source,title in [(axes[0],table,'A. Prespecified finite heat shift'),
                                    (axes[1],refined,'B. Exploratory range check')]:
                part=source[source.fit_type=='finite_nonnegative'].pivot(index='weak',columns='t',values='test_ratio').loc[names,ts]
                im=ax.imshow(part.to_numpy(),vmin=0,vmax=1.05,cmap='YlOrRd',aspect='auto')
                ax.set_xticks(range(len(ts)),[f'{t:g}' for t in ts]);ax.set_yticks(range(len(names)),names)
                ax.set(xlabel='Flow time t (noise -> data)',title=title)
                for i in range(len(names)):
                    for j in range(len(ts)):
                        value=part.iloc[i,j]
                        ax.text(j,i,f'{value:.3f}',ha='center',va='center',color='white' if value>.65 else 'black')
            fig.colorbar(im,ax=axes,label='Held-out residual / original strong-weak gap energy',shrink=.85)
            fig.savefig(OUT/'neural_range_check.png');fig.savefig(OUT/'neural_range_check.pdf');plt.close(fig)


def workbook_and_validation():
    # Reuse the isolated dependency directory from the preceding audit if needed.
    try:
        import openpyxl
    except ModuleNotFoundError:
        sys.path.insert(0,'/tmp/eqvae_ag_audit_dependencies')
        import openpyxl
    tables=sorted(OUT.glob('*.csv'))
    wb=openpyxl.Workbook();ws=wb.active;ws.title='README'
    for row in [
        ['AG / IG smoothing audit','2026-09-11'],
        ['Purpose','Exact controls and CPU field consistency; not a new FID result'],
        ['Neural split','32 validation images: 16 fit, 16 held out; same images across 5 times'],
        ['Main denominator','Energy of strong-weak score gap'],
        ['All figures','Source rows in the named CSV sheets; formulas in report script'],
        ['Raw observations','/home/zhoushunyu/data/eqvae/imagenet_sit_flow/ag_ig_smoothing_cpu_20260911'],
    ]:ws.append(row)
    for path in tables:
        df=pd.read_csv(path)
        sheet=wb.create_sheet(path.stem[:31]);sheet.append(df.columns.tolist())
        for row in df.itertuples(index=False,name=None):sheet.append(row)
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    wb.save(OUT/'ag_ig_smoothing_support.xlsx')
    checked=openpyxl.load_workbook(OUT/'ag_ig_smoothing_support.xlsx',read_only=True,data_only=True)
    counts={p.stem:len(pd.read_csv(p)) for p in tables}
    for p in tables:assert checked[p.stem[:31]].max_row==counts[p.stem]+1
    if (OUT/'neural_summary.csv').exists():
        summary=pd.read_csv(OUT/'neural_summary.csv')
        per=pd.read_csv(OUT/'neural_per_image.csv');per=per[per.split=='test']
        grouped=per.groupby(['weak','t','fit_type'])[['squared_residual','squared_gap']].sum()
        for row in summary.itertuples():
            g=grouped.loc[(row.weak,row.t,row.fit_type)]
            assert abs(g.squared_residual/g.squared_gap-row.test_ratio)<1e-10
        curves=pd.read_csv(OUT/'neural_fit_curves.csv')
        assert np.max(abs(curves[curves.delta==0].fit_ratio-1))<1e-12
        assert np.max(abs(curves[curves.delta==0].test_ratio-1))<1e-12
        assert len(summary)==100 and len(per)==1600
        assert np.isfinite(summary.select_dtypes('number')).all().all()
        if (OUT/'refined_summary.csv').exists():
            refined=pd.read_csv(OUT/'refined_summary.csv')
            rp=pd.read_csv(OUT/'refined_per_image.csv');rp=rp[rp.split=='test']
            grouped=rp.groupby(['weak','t','fit_type'])[['squared_residual','squared_gap']].sum()
            for row in refined.itertuples():
                g=grouped.loc[(row.weak,row.t,row.fit_type)]
                assert abs(g.squared_residual/g.squared_gap-row.test_ratio)<1e-10
            assert len(refined)==50 and len(rp)==800
    raw_checks={}
    raw=Path('/home/zhoushunyu/data/eqvae/imagenet_sit_flow/ag_ig_smoothing_cpu_20260911')
    if (OUT/'refined_summary.csv').exists():
        original=json.loads((raw/'manifest.json').read_text())
        stage=raw/'range_check_20260912'
        amended=json.loads((stage/'manifest.json').read_text())
        inputs=np.load(raw/'inputs.npz')
        assert len(set(inputs['ids'].tolist()))==32
        assert original['completed'] and amended['completed']
        assert not original['cuda_initialized'] and not amended['cuda_initialized']
        checked_hashes=0
        for directory,manifest in [(raw,original),(stage,amended)]:
            for filename,sha in manifest['observations'].items():
                assert hashlib.sha256((directory/filename).read_bytes()).hexdigest()==sha
                checked_hashes+=1
        assert hashlib.sha256((raw/'inputs.npz').read_bytes()).hexdigest()==original['inputs_sha256']
        assert hashlib.sha256((ROOT/'experiments/audit_ag_ig_smoothing_20260911.py').read_bytes()).hexdigest()==original['script_sha256']
        assert hashlib.sha256((ROOT/'experiments/refine_ag_ig_smoothing_20260912.py').read_bytes()).hexdigest()==amended['script_sha256']
        maximum_ratio_error=0.;maximum_scaled_ratio_error=0.;maximum_state_error=0.;maximum_native_parity=0.
        all_curves=pd.read_csv(OUT/'refined_fit_curves.csv')
        for ti,t in enumerate(original['times']):
            base=np.load(raw/f'time_{ti}.npz');extra=np.load(stage/f'time_{ti}.npz')
            state=(1-t)*inputs['noise']+t*inputs['clean']
            maximum_state_error=max(maximum_state_error,float(np.max(abs(state-base['z']))))
            maximum_native_parity=max(maximum_native_parity,float(base['parity_max']))
            # Reconstruct all candidate ratios directly from registered fields,
            # using a separate scalar reduction, rather than trusting row summaries.
            for row in all_curves[all_curves.t==t].itertuples():
                key=f'shift_{row.delta:g}'
                shifted=extra[key] if key in extra else base[key]
                w=base[f'weak_{row.weak}'].astype(np.float64)
                diff=w-base['strong']
                error=w-shifted
                for sl,stored in [(slice(0,16),row.fit_ratio),(slice(16,32),row.test_ratio)]:
                    ratio=float(np.vdot(error[sl],error[sl])/np.vdot(diff[sl],diff[sl]))
                    maximum_ratio_error=max(maximum_ratio_error,abs(ratio-stored))
                    maximum_scaled_ratio_error=max(maximum_scaled_ratio_error,abs(ratio-stored)/max(1.,abs(stored)))
            base.close();extra.close()
        # Some deliberately wrong shifts have ratios above 200,000; their
        # independent reductions differ at machine precision in relative terms.
        assert maximum_scaled_ratio_error<1e-12 and maximum_state_error==0
        sources=json.loads((ROOT/'readings/ag_ig_smoothing_20260911/source_manifest.json').read_text())
        for source in sources:
            assert hashlib.sha256((ROOT/source['path']).read_bytes()).hexdigest()==source['sha256']
        raw_checks=dict(registered_observation_hashes_verified=checked_hashes,source_hashes_verified=len(sources),
                        all_candidate_ratio_max_error=maximum_ratio_error,all_candidate_ratio_max_scaled_error=maximum_scaled_ratio_error,
                        teacher_state_max_error=maximum_state_error,
                        native_forward_parity_max_error=maximum_native_parity,unique_ids=32,cuda_initialized=False)
    valid=dict(passed=True,csv_rows=counts,workbook_sheets=checked.sheetnames,raw_checks=raw_checks,
               hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [*tables,OUT/'ag_ig_smoothing_support.xlsx']})
    (OUT/'artifact_validation.json').write_text(json.dumps(valid,indent=2)+'\n')


if __name__=='__main__':
    print(json.dumps(exact_controls(),indent=2))
    figures()
    workbook_and_validation()
