"""Finalize only after samplers, source training and bounded screens complete."""
import os
from pathlib import Path
import time
import zipfile
import numpy as np
import openpyxl
import torch
from experiments.guidance_pasted_20260912 import common as c
from . import mixture as m,audit,report

STOPS={
 'sit_refined_priority_20260911':'c8227c2a9aa9ccf49003bdb1652a4a6cce870fb79ae8066a7a475811e3bb3635',
 'sit_apg_mechanism_extension_20260911':'c8227c2a9aa9ccf49003bdb1652a4a6cce870fb79ae8066a7a475811e3bb3635',
 'sit_broad_resume_20260912':'c8227c2a9aa9ccf49003bdb1652a4a6cce870fb79ae8066a7a475811e3bb3635',
 'sit_control_output_50ideas_20260910':'b43aef680652b33cc5a199cecb355a717ae50af5e058cecaa83d4f91aa180ce1',
 'sit_reference_aggregation_20260912':'f17cc7a0cb994382cf856a190d0603caa6d452cbd62211c06075a16d60b7a09d'}


def alive(pid):
    try:os.kill(pid,0);return True
    except ProcessLookupError:return False


def source_audit(model):
    root=m.ROOT/model/'cfg_source_data';req=root/'request.json';request=c.read(req)
    for group in ('sources','assets','inputs'):
        for p,h in request[group].items():assert c.sha(p)==h,p
    bank=np.load(root/'inputs/first.npy',mmap_mode='r');labels=np.load(root/'inputs/labels.npy')
    summary=c.read(root/'training_complete.json');coverage={0:[],1:[]};seconds=0.
    initial={0:{},1:{}};q=96 if model=='sit_small' else 100
    for r in summary['records']:
        p=Path(r['file']);assert c.sha(p)==r['sha256'];meta=c.read(p.with_suffix('.json'))
        with np.load(p) as d:
            ids=d['ids'];label=int(d['source']);coverage[label].extend(ids)
            assert np.array_equal(ids,np.arange(ids[0],ids[0]+len(ids)))
            assert str(d['noise_sha256'])==c.array_sha(bank[ids[0]:ids[0]+len(ids)])
            np.testing.assert_array_equal(d['labels'],labels[ids])
            assert str(d['request_sha256'])==c.sha(req)
            assert d['features'].shape[1]==q and np.isfinite(d['features']).all()
            for i,idx in enumerate(ids):initial[label][int(idx)]=d['features'][i,0]
            seconds+=float(d['seconds'])
        assert meta['counts']==dict(full=q,prefix=q if label==0 else 0)
    for v in coverage.values():assert sorted(v)==list(range(800))
    for i in range(800):np.testing.assert_array_equal(initial[0][i],initial[1][i])
    assert abs(seconds-summary['source_seconds'])<1e-6
    assert c.sha(root/'head.pt')==summary['head_sha256']
    r=dict(passed=True,source_files=len(summary['records']),seeds_per_source=800,queries_per_seed=q,
        exact_initial_features_for_all_seeds=True,noise_labels_and_frozen_request_verified=True,
        full_and_prefix_counts_verified=True,source_seconds=seconds,request_sha256=c.sha(req))
    c.atomic(root/'audit.json',r);return r


def main():
    torch.set_num_threads(4);m.configure()
    for status in ('endpoint_status.json','ig_mixture_status.json',
                   'sit_small_cfg_pipeline_status.json','raev2_cfg_pipeline_status.json'):
        p=m.ROOT/status
        while not p.exists() or c.read(p)['phase']!='complete':
            if p.exists() and not alive(c.read(p)['pid']):raise RuntimeError(f'Pipeline exited: {p}')
            time.sleep(8)
        pid=c.read(p)['pid']
        while alive(pid):time.sleep(2)
    for model in c.MODELS:
        while not (m.ROOT/model/'inference_benchmark.json').exists():time.sleep(5)
    audits=[];sa=[]
    for model in c.MODELS:
        for stage in ('endpoint_sources',m.stage('ig'),m.stage('cfg')):
            audit.run(model,stage)
            audits.extend(c.read(m.ROOT/model/stage/'audit.json')['records'])
        sa.append(source_audit(model))
        tr=m.ROOT/model/'input_local_training'
        assert c.read(tr/'summary.json')['complete']
        assert not (m.ROOT/model/'input_local_screen_400').exists()
        req=c.read(tr/'request.json')
        for group in ('sources','assets','data'):
            for p,h in req[group].items():assert c.sha(p)==h,p
        assert c.read(tr/'generation_deferred.json')['generation_started'] is False
    stops={}
    for name,h in STOPS.items():
        p=c.EXPS/name/'STOP_AFTER_CURRENT'
        assert c.sha(p)==h,(p,c.sha(p))
        stops[str(p)]=h
    report.main()
    with zipfile.ZipFile(report.OUT/'source_data.xlsx') as z:assert z.testzip() is None
    workbook=openpyxl.load_workbook(report.OUT/'source_data.xlsx',read_only=True,data_only=False)
    assert workbook['quality'].max_row==25 and workbook['moments'].max_row==5
    assert workbook['inference_timing'].max_row==5
    workbook.close()
    tables=report.frames();eligible=tables['quality'][tables['quality'].passes_fixed_screen]
    final=dict(passed=True,quality_arms=24,source_endpoint_arms=6,
        quality_images=9600,source_endpoint_images=6000,new_cfg_partial_source_paths=3200,
        total_full_generated_images=15600,source_files=sum(r['source_files'] for r in sa),
        max_cached_feature_fid_recalculation_error=max(r['absolute_error'] for r in audits),
        independent_feature_reextraction=False,source_audits=sa,old_stops_unchanged=stops,
        input_local_quality_arms=0,source_workbook_valid=True,
        quality_sample_decode_gpu_batch_seconds=float(tables['quality'].seconds.sum()),
        endpoint_sample_decode_gpu_batch_seconds=float(tables['endpoint_sources'].seconds.sum()),
        cfg_source_gpu_batch_seconds=float(tables['cfg_classifier'].source_seconds.sum()),
        eligible_for_fresh_1k=eligible[['model','track']].to_dict(orient='records'),
        benchmark_native_pixel_parity=True,long_term_goal_achieved=False,
        visual_inspection_pending=True)
    c.atomic(report.OUT/'final_verification.json',final)
    c.atomic(m.ROOT/'round_status.json',dict(phase='round_complete',main_hypothesis='shared contamination retained',
        quality_arms=24,eligible_for_fresh_1k=final['eligible_for_fresh_1k'],
        input_local_generation_deferred=True,old_large_queues_stopped=True))
    # Refresh the manifest after adding final verification.
    c.atomic(report.OUT/'artifact_manifest.json',dict(
        files={str(p.relative_to(c.WORK)):c.sha(p) for p in sorted(report.OUT.iterdir()) if p.name!='artifact_manifest.json'},
        report_sha256=c.sha(report.REPORT),generator_sha256=c.sha(Path(report.__file__))))
    print('Round finalized',final,flush=True)


if __name__=='__main__':main()
