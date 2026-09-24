"""Small shared scene model for editable SVG/PPTX scientific diagrams."""
from pathlib import Path
import base64, io, json, math, os, random
from xml.sax.saxutils import escape
ROOT=Path(__file__).resolve().parent
config=ROOT/'fonts'/'fontconfig.xml'
config.write_text(f'<fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include><dir>{ROOT/"fonts"}</dir></fontconfig>')
os.environ['FONTCONFIG_FILE']=str(config)
import cairosvg
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
W,H=1080,720
NAVY,BLUE,RED='#0E2841','#3072D5','#D74732'
SCENE=[]
FONT_FAMILY='Comic Sans MS'
FONT_FILES={False:'Comic.TTF',True:'Comicbd.TTF'}

def add(kind, name, **kwargs):
    SCENE.append(dict(kind=kind, name=name, **kwargs))

def path(name, commands, fill=None, stroke=None, width=1, opacity=1):
    # Icons can reuse command templates. Give each shape independent geometry
    # so subsequent layout translations move every path exactly once.
    add('path', name, commands=[list(c) for c in commands], fill=fill, stroke=stroke,
        width=width, opacity=opacity)

def ellipse(name, cx, cy, rx, ry, fill, opacity=1, stroke=None, width=.5):
    add('ellipse', name, cx=cx, cy=cy, rx=rx, ry=ry, fill=fill,
        opacity=opacity, stroke=stroke, width=width)

def text(name, x, y, w, h, content, size=13.7, fill=NAVY, bold=False,
         font=FONT_FAMILY, align='center'):
    # Rich runs use native text with an explicit subscript baseline.
    runs = content if isinstance(content, list) else [dict(text=content)]
    add('text', name, x=x, y=y, w=w, h=h, runs=runs, size=size,
        fill=fill, bold=bold, font=font, align=align)

def snow(name, cx, cy, r=10):
    cmds=[]
    for a in range(0,360,60):
        c,s=math.cos(math.radians(a)),math.sin(math.radians(a))
        def p(u,v):return [cx+c*u-s*v,cy+s*u+c*v]
        cmds += [['M',*p(0,0)],['L',*p(r,0)]]
        for u in (r*.53,r*.76):
            for sign in (-1,1):
                cmds += [['M',*p(u,0)],['L',*p(u-r*.23,sign*r*.23)]]
    path(name+'/white outline',cmds,stroke='#FFFFFF',width=2.5)
    path(name+'/blue crystal',cmds,stroke='#83B8DC',width=1)

def flame(name,cx,cy,scale=1):
    outer=[['M',0,12],['C',-12,8,-8,-2,-5,-5],['C',-6,1,-1,2,-2,-4],
           ['C',-4,-8,3,-10,0,-15],['C',11,-10,6,-3,9,-6],
           ['C',15,1,12,10,5,12],['C',3,13,1,13,0,12],['Z']]
    inner=[['M',1,11],['C',-4,9,-4,3,-1,0],['C',-1,5,3,5,2,-2],
           ['C',8,2,8,9,3,11],['Z']]
    for tag,cmds,fill in [('red',outer,'#FF182E'),('gold',inner,'#FFD546')]:
        q=[]
        for c in cmds:
            q.append([c[0]]+[v*scale+(cx if j%2==0 else cy) for j,v in enumerate(c[1:])])
        path(name+'/'+tag,q,fill=fill)

def cloud(name,cx,cy,rx,ry,fill,seed,n=75):
    for factor,op in [(1.35,.045),(1.05,.07),(.75,.10)]:
        ellipse(name+f'/contour {factor}',cx,cy,rx*factor,ry*factor,fill,op,fill,.35)
    rng=random.Random(seed)
    for j in range(n):
        while True:
            a,b=rng.gauss(0,.37),rng.gauss(0,.37)
            if a*a+b*b<1.3:break
        r=rng.uniform(.65,1.25)
        ellipse(name+f'/point {j:03d}',cx+a*rx,cy+b*ry,r,r,fill,rng.uniform(.5,.88))

def el(tag, **attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items():e.set(k,str(v))
    return e

def paint(sh,s):
    # Override the default Office shape style's shadow; paper diagrams are flat.
    # LibreOffice also needs the theme effect reference disabled explicitly.
    for effect_ref in sh._element.xpath('./p:style/a:effectRef'):
        effect_ref.set('idx','0')
    sp=sh._element.spPr
    for child in list(sp):
        if child.tag.endswith('}effectLst') or child.tag.endswith('}effectDag'):sp.remove(child)
    sp.append(el('a:effectLst'))
    if s['fill']:
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(s['fill'][1:])
        if s.get('opacity',1)<1:
            sh.fill._xPr.solidFill[0].append(el('a:alpha',val=round(s['opacity']*100000)))
    else:sh.fill.background()
    if s['stroke']:
        sh.line.color.rgb=RGBColor.from_string(s['stroke'][1:]);sh.line.width=Pt(s['width'])
        if s.get('opacity',1)<1:
            sh.line._get_or_add_ln().find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr').append(el('a:alpha',val=round(s['opacity']*100000)))
    else:sh.line.fill.background()


def rect(name,x,y,w,h,fill,stroke=None,r=7,width=1):
    k=.55228475
    path(name,[['M',x+r,y],['L',x+w-r,y],
        ['C',x+w-r+k*r,y,x+w,y+r-k*r,x+w,y+r],['L',x+w,y+h-r],
        ['C',x+w,y+h-r+k*r,x+w-r+k*r,y+h,x+w-r,y+h],['L',x+r,y+h],
        ['C',x+r-k*r,y+h,x,y+h-r+k*r,x,y+h-r],['L',x,y+r],
        ['C',x,y+r-k*r,x+r-k*r,y,x+r,y],['Z']],fill,stroke,width)


def line(name,points,stroke=NAVY,width=1.5,dash=False,end=False):
    path(name,[['M',*points[0]]]+[['L',*p] for p in points[1:]],stroke=stroke,width=width)
    if dash:SCENE[-1]['dash']=[5,3]
    if end:
        x,y=points[-1];px,py=points[-2];a=math.atan2(y-py,x-px)
        length=6;half=2.7
        q1=[x-length*math.cos(a)+half*math.sin(a),y-length*math.sin(a)-half*math.cos(a)]
        q2=[x-length*math.cos(a)-half*math.sin(a),y-length*math.sin(a)+half*math.cos(a)]
        path(name+' head',[['M',x,y],['L',*q1],['L',*q2],['Z']],fill=stroke)


def measure(s):
    widths=[]
    for r in s['runs']:
        f=FONT_FILES[s['bold']]
        size=s['size']*(.73 if r.get('sub') or r.get('sup') else 1)
        font=ImageFont.truetype(str(ROOT/'fonts'/f),round(size*100))
        widths.append(font.getlength(r['text'])/100)
    return widths


def svg(scene,title):
    fonts=[]
    for family,files in [(FONT_FAMILY,[FONT_FILES[False],FONT_FILES[True]])]:
        for weight,f in zip([400,700],files):
            data=base64.b64encode((ROOT/'fonts'/f).read_bytes()).decode()
            fonts.append(f"@font-face{{font-family:'{family}';font-weight:{weight};src:url(data:font/ttf;base64,{data})}}")
    out=[f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="{W}pt" height="{H}pt" viewBox="0 0 {W} {H}">',
         f'<title>{escape(title)}</title>',
         '<desc>Classifier-guided learning of weak predictions and signed guidance coefficients. Style adapted from AdvFD Figure 3 (arXiv:2608.11205v1). All diagrams and scatter points are schematic, not empirical results.</desc>',
         '<defs><style>'+''.join(fonts)+'</style></defs>',f'<rect width="{W}" height="{H}" fill="white"/>']
    for i,s in enumerate(scene):
        out.append(f'<g id="object-{i:04d}" data-name="{escape(s["name"])}"><title>{escape(s["name"])}</title>')
        if s['kind'] in ('path','ellipse'):
            attr=f'fill="{s["fill"] or "none"}" stroke="{s["stroke"] or "none"}" stroke-width="{s["width"]}" opacity="{s["opacity"]}" stroke-linecap="round" stroke-linejoin="round"'
            if s.get('dash'):attr+=' stroke-dasharray="'+' '.join(map(str,s['dash']))+'"'
            if s['kind']=='path':
                d=' '.join(c[0]+' '+' '.join(f'{v:.4f}' for v in c[1:]) for c in s['commands'])
                out.append(f'<path d="{d}" {attr}/>')
            else:out.append('<ellipse '+' '.join(f'{k}="{s[k]}"' for k in ('cx','cy','rx','ry'))+' '+attr+'/>')
        elif s['kind']=='image':
            out.append(f'<image x="{s["x"]}" y="{s["y"]}" width="{s["w"]}" height="{s["h"]}" xlink:href="data:image/png;base64,{s["data"]}"/>')
        else:
            widths=measure(s)
            x=s['x']+((s['w']-sum(widths))/2 if s['align']=='center' else 0)
            y=s['y']+s['h']/2+s['size']*.32
            if s.get('rotation'):
                out.append(f'<g transform="rotate({s["rotation"]},{s["x"]+s["w"]/2},{s["y"]+s["h"]/2})">')
            for r,w in zip(s['runs'],widths):
                shift=.24*s['size'] if r.get('sub') else (-.35*s['size'] if r.get('sup') else 0)
                size=s['size']*(.73 if r.get('sub') or r.get('sup') else 1)
                italic=' font-style="italic"' if r.get('italic') else ''
                out.append(f'<text x="{x:.4f}" y="{y+shift:.4f}" font-family="{r.get("font",s["font"])}" font-size="{size}" font-weight="{700 if s["bold"] else 400}" fill="{s["fill"]}"{italic} xml:space="preserve">{escape(r["text"])}</text>')
                x+=w
            if s.get('rotation'):out.append('</g>')
        out.append('</g>')
    out.append('</svg>')
    return ''.join(out)


def add_slide(prs,scene,title,notes):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.name=title
    group=None;last_group=None
    for s in scene:
        group_name=s['name'].split('/')[0]
        # Consecutive grouping preserves exactly the scene's z-order.
        if group_name!=last_group:
            group=slide.shapes.add_group_shape();group.name=group_name;last_group=group_name
            group.shapes.turbo_add_enabled=True
        target=group.shapes
        if s['kind']=='ellipse':
            sh=target.add_shape(MSO_SHAPE.OVAL,Pt(s['cx']-s['rx']),Pt(s['cy']-s['ry']),Pt(2*s['rx']),Pt(2*s['ry']))
            paint(sh,s)
        elif s['kind']=='path':
            coords=[(c[i],c[i+1]) for c in s['commands'] for i in range(1,len(c),2)]
            x0=min(p[0] for p in coords);y0=min(p[1] for p in coords)
            w=max(max(p[0] for p in coords)-x0,.01);h=max(max(p[1] for p in coords)-y0,.01)
            sh=target.add_shape(MSO_SHAPE.RECTANGLE,Pt(x0),Pt(y0),Pt(w),Pt(h))
            sp=sh._element.spPr;sp.remove(sp.prstGeom)
            geom=el('a:custGeom')
            for tag in ['a:avLst','a:gdLst','a:ahLst','a:cxnLst']:geom.append(el(tag))
            geom.append(el('a:rect',l='0',t='0',r='r',b='b'))
            pl=el('a:pathLst');pa=el('a:path',w=round(w*1000),h=round(h*1000))
            for c in s['commands']:
                if c[0]=='Z':pa.append(el('a:close'));continue
                cmd=el({'M':'a:moveTo','L':'a:lnTo','C':'a:cubicBezTo'}[c[0]])
                for i in range(1,len(c),2):cmd.append(el('a:pt',x=round((c[i]-x0)*1000),y=round((c[i+1]-y0)*1000)))
                pa.append(cmd)
            pl.append(pa);geom.append(pl);sp.insert(1,geom);paint(sh,s)
            if s.get('dash'):sh.line._get_or_add_ln().append(el('a:prstDash',val='dash'))
        elif s['kind']=='image':
            sh=target.add_picture(io.BytesIO(base64.b64decode(s['data'])),Pt(s['x']),Pt(s['y']),Pt(s['w']),Pt(s['h']))
        else:
            sh=target.add_textbox(Pt(s['x']),Pt(s['y']),Pt(s['w']),Pt(s['h']))
            tf=sh.text_frame;tf.clear();tf.word_wrap=False
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            tf.vertical_anchor=MSO_ANCHOR.MIDDLE
            p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER if s['align']=='center' else PP_ALIGN.LEFT
            p.font.name=FONT_FAMILY
            p.space_before=p.space_after=Pt(0)
            for r in s['runs']:
                rr=p.add_run();rr.text=r['text'];f=rr.font
                # Office scales a run automatically when baseline is nonzero.
                # Keeping the parent size avoids shrinking subscripts twice.
                f.name=r.get('font',s['font']);f.size=Pt(s['size'])
                f.bold=s['bold'];f.italic=r.get('italic',False);f.color.rgb=RGBColor.from_string(s['fill'][1:])
                rpr=rr._r.get_or_add_rPr()
                for tag in ['a:ea','a:cs']:
                    rpr.append(el(tag,typeface=FONT_FAMILY))
                if r.get('sub'):rr._r.get_or_add_rPr().set('baseline','-25000')
                if r.get('sup'):rr._r.get_or_add_rPr().set('baseline','35000')
            if s.get('rotation'):sh.rotation=s['rotation']%360
        sh.name=s['name']
    slide.notes_slide.notes_text_frame.text=notes
    return slide


def export_deck(slides,path):
    from embedded_fonts import embed_family
    prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
    prs.core_properties.title='Adversarial learning of self-guidance — method overview'
    prs.core_properties.subject='Source-backed editable method figures for joint and schedule-only learning'
    prs.core_properties.comments='Visual style adapted from AdvFD Figure 3; content based on local classifier_guidance implementation.'
    for scene,title,notes in slides:add_slide(prs,scene,title,notes)
    embed_family(prs,FONT_FAMILY,ROOT/'fonts'/FONT_FILES[False],ROOT/'fonts'/FONT_FILES[True])
    prs.save(path)
