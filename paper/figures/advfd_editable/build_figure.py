"""Rebuild AdvFD Figure 3 as native, editable PowerPoint and SVG objects.

Input: the vector figure from arXiv:2608.11205v1, images/method_new.pdf.
Run with: python build_figure.py (python-pptx, pymupdf, cairosvg, Pillow).
"""
from pathlib import Path
import base64
import io
import json
import math
import os
import random
import re
import unicodedata
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent
# Register the supplied open font for Cairo without modifying system fonts.
FONTCONFIG = ROOT / 'fonts' / 'fontconfig.xml'
FONTCONFIG.write_text(f'''<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd">
<fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>
<dir>{ROOT / 'fonts'}</dir></fontconfig>''')
os.environ['FONTCONFIG_FILE'] = str(FONTCONFIG)

import pymupdf as fitz
import cairosvg
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

W, H, DX, DY = 837, 478, 8, 4
NAVY, BLUE, RED = '#0E2841', '#3072D5', '#D74732'
SCENE = []
DOC = fitz.open(ROOT / 'reference' / 'advfd_figure3_original.pdf')
PAGE = DOC[0]


def color(v):
    return '#' + ''.join(f'{round(c * 255):02X}' for c in v) if v else None


def add(kind, name, **kwargs):
    SCENE.append(dict(kind=kind, name=name, **kwargs))


def path(name, commands, fill=None, stroke=None, width=1, opacity=1):
    add('path', name, commands=commands, fill=fill, stroke=stroke,
        width=width, opacity=opacity)


def ellipse(name, cx, cy, rx, ry, fill, opacity=1, stroke=None, width=.5):
    add('ellipse', name, cx=cx, cy=cy, rx=rx, ry=ry, fill=fill,
        opacity=opacity, stroke=stroke, width=width)


def text(name, x, y, w, h, content, size=13.7, fill=NAVY, bold=False,
         font='Comic Neue', align='center'):
    # Rich runs use native text with an explicit subscript baseline.
    runs = content if isinstance(content, list) else [dict(text=content)]
    add('text', name, x=x, y=y, w=w, h=h, runs=runs, size=size,
        fill=fill, bold=bold, font=font, align=align)


def pdf_geometry():
    omitted = {0, 3, 4, 11, 13, 19, 48, 50}  # page background, equation parentheses
    names = {2:'G-step/objective background', 10:'D-step/static loss background',
             12:'D-step/adversarial loss background',18:'D-step/objective background',
             23:'D-step/frozen generator',41:'D-step/static representation',
             42:'D-step/trainable adverse representation',47:'G-step/static loss background',
             49:'G-step/adversarial loss background',54:'G-step/trainable generator',
             72:'G-step/static representation',73:'G-step/adverse representation'}
    for i, d in enumerate(PAGE.get_drawings()):
        if i in omitted:
            continue
        commands, current = [], None
        for item in d['items']:
            op = item[0]
            if op in ('l', 'c'):
                start = tuple(item[1])
                if current != start:
                    commands.append(['M', *start])
                if op == 'l':
                    commands.append(['L', *item[2]])
                else:
                    commands.append(['C', *item[2], *item[3], *item[4]])
                current = tuple(item[-1])
            elif op == 're':
                r = item[1]
                commands.extend([['M',r.x0,r.y0],['L',r.x1,r.y0],
                    ['L',r.x1,r.y1],['L',r.x0,r.y1],['Z']])
                current = None
        if d.get('closePath'):
            commands.append(['Z'])
        path(names.get(i, f'Original geometry/{i:02d}'), commands,
             color(d['fill']), color(d['color']), d['width'] or 1)


def snow(name, cx, cy, r=10):
    cmds=[]
    for a in range(0,360,60):
        c,s=math.cos(math.radians(a)),math.sin(math.radians(a))
        def p(u,v):return [cx+c*u-s*v,cy+s*u+c*v]
        cmds += [['M',*p(0,0)],['L',*p(r,0)]]
        for u in (r*.53,r*.76):
            for sign in (-1,1):
                cmds += [['M',*p(u,0)],['L',*p(u-r*.23,sign*r*.23)]]
    path(name+'/white outline',cmds,stroke='#FFFFFF',width=2.5)
    path(name+'/blue crystal',cmds,stroke='#83B8DC',width=1)


def flame(name,cx,cy,scale=1):
    outer=[['M',0,12],['C',-12,8,-8,-2,-5,-5],['C',-6,1,-1,2,-2,-4],
           ['C',-4,-8,3,-10,0,-15],['C',11,-10,6,-3,9,-6],
           ['C',15,1,12,10,5,12],['C',3,13,1,13,0,12],['Z']]
    inner=[['M',1,11],['C',-4,9,-4,3,-1,0],['C',-1,5,3,5,2,-2],
           ['C',8,2,8,9,3,11],['Z']]
    for tag,cmds,fill in [('red',outer,'#FF182E'),('gold',inner,'#FFD546')]:
        q=[]
        for c in cmds:
            q.append([c[0]]+[v*scale+(cx if j%2==0 else cy) for j,v in enumerate(c[1:])])
        path(name+'/'+tag,q,fill=fill)


def cloud(name,cx,cy,rx,ry,fill,seed,n=75):
    for factor,op in [(1.35,.045),(1.05,.07),(.75,.10)]:
        ellipse(name+f'/contour {factor}',cx,cy,rx*factor,ry*factor,fill,op,fill,.35)
    rng=random.Random(seed)
    for j in range(n):
        while True:
            a,b=rng.gauss(0,.37),rng.gauss(0,.37)
            if a*a+b*b<1.3:break
        r=rng.uniform(.65,1.25)
        ellipse(name+f'/point {j:03d}',cx+a*rx,cy+b*ry,r,r,fill,rng.uniform(.5,.88))


def equation(name,x,y,w,which,bold=False):
    r=[]
    def t(s,sub=False): r.append(dict(text=s,sub=sub,italic=True))
    if which=='g':
        t('min');t('θ',True);t(' D');t('static',True);t('(p, q');t('θ',True)
        t(') + λ');t('adv',True);t(' ∗ D');t('adv',True);t('(p, q');t('θ',True);t(')')
    elif which=='d':
        t('max');t('ψ',True);t(' D');t('adv',True);t('(p, q');t('θ',True);t(')')
    else:
        t('D');t(which,True);t('(p, q');t('θ',True);t(')')
    text(name,x,y,w,26,r,12.5 if which=='g' else 13.7,
         '#C14F15' if bold else NAVY,bold,'DejaVu Serif')


def build_scene():
    pdf_geometry()
    # The same original noise tile remains a replaceable picture object.
    tile=DOC.extract_image(27)['image']
    add('image','G-step/noise texture',x=94.24,y=14.449,w=72.72,h=42.491,data=base64.b64encode(tile).decode())
    add('image','D-step/noise texture',x=524.08,y=14.449,w=73.44,h=42.491,data=base64.b64encode(tile).decode())
    for off,side in [(0,'G-step'),(430.68,'D-step')]:
        text(side+'/Noise',94+off,59,74,21,'Noise',14.5,fill='#000000')
        text(side+'/Generator',211+off,24,115,25,[dict(text='Generator '),dict(text='G',font='DejaVu Serif'),dict(text='θ',sub=True,font='DejaVu Serif')],14.4)
        text(side+'/Real Data',84+off,107,94,20,'Real Data',12.8,fill='#E1706B')
        text(side+'/Fake Batch',215+off,112,110,20,'Fake Batch',12.8,fill='#DFB882',align='left')
        text(side+'/Static Repr',62+off,178,142,27,[dict(text='Static Repr. '),dict(text='φ',italic=True,font='DejaVu Serif')],14.4)
        text(side+'/Adverse Repr',231+off,178,144,27,[dict(text='Adverse Repr. '),dict(text='ψ',italic=True,font='DejaVu Serif')],14.4)
        equation(side+'/Static FD',58+off,233,144,'static')
        equation(side+'/Adverse FD',231+off,233,144,'adv')
    equation('G-step/min objective',63,293,313,'g',True)
    equation('D-step/max objective',491,293,313,'d',True)
    for name,cx,cy in [('G-step/static frozen',67,173),('G-step/adverse frozen',235,173),
                      ('D-step/static frozen',498,173),('D-step/generator frozen',647,14)]:
        snow(name,cx,cy)
    flame('G-step/generator training',215,14,.8)
    flame('D-step/adverse training',666,174,.8)
    # Draw contours and scatter points as individual editable objects.
    cloud('G-before/real',128,410,24,22,RED,9)
    cloud('G-before/fake',171,365,24,27,BLUE,19)
    cloud('G-after/real',328,383,32,32,RED,29,90)
    cloud('G-after/fake',338,381,32,35,BLUE,39,90)
    cloud('D-before/real',536,392,27,20,RED,49)
    cloud('D-before/fake',566,389,26,20,BLUE,59)
    cloud('D-after/real',707,390,23,18,RED,69)
    cloud('D-after/fake',790,389,23,18,BLUE,79)
    # Restore explanatory arrows above the point clouds.
    original_arrows=[s.copy() for s in SCENE if s['name'] in ['Original geometry/06','Original geometry/07','Original geometry/21','Original geometry/22']]
    for s in original_arrows:
        SCENE.remove(next(v for v in SCENE if v['name']==s['name']))
        SCENE.append(s)
    text('G-before/real label',0,351,92,24,'Real Samples',14.3,'#D02B10',align='left')
    text('G-before/fake label',18,322,104,24,'Fake Samples',14.3,'#679FEE',align='left')
    text('D-before/real label',402,361,92,24,'Real Samples',14.3,'#D02B10',align='left')
    text('D-before/fake label',420,332,103,24,'Fake Samples',14.3,'#679FEE',align='left')
    text('G-step/update label',208,362,75,24,[dict(text='Update '),dict(text='G',font='DejaVu Serif'),dict(text='θ',sub=True,font='DejaVu Serif')],14.3,fill='#000000')
    text('G-step/action label',208,395,75,24,'Pull Closer',14.3,fill='#000000')
    text('D-step/update label',600,362,76,24,[dict(text='Update '),dict(text='ψ',italic=True,font='DejaVu Serif')],14.3,fill='#000000')
    text('D-step/action label',600,395,76,24,'Push Apart',14.3,fill='#000000')
    text('G-step/panel title',143,439,182,30,'(a) G-step Training',16.5,'#000000',True)
    text('D-step/panel title',564,439,182,30,'(b) D-step Training',16.5,'#000000',True)


def svg_render():
    fonts=[]
    for weight,f in [(400,'ComicNeue-Regular.ttf'),(700,'ComicNeue-Bold.ttf')]:
        data=base64.b64encode((ROOT/'fonts'/f).read_bytes()).decode()
        fonts.append(f"@font-face{{font-family:'Comic Neue';font-weight:{weight};src:url(data:font/ttf;base64,{data})}}")
    out=[f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="{W}pt" height="{H}pt" viewBox="0 0 {W} {H}">',
         '<title>AdvFD Figure 3 — editable reconstruction</title>',
         '<desc>Reconstructed from arXiv:2608.11205v1, Figure 3. Original layout and paths; editable text and recreated illustrative scatter points. This is a reference figure, not our method.</desc>',
         '<defs><style>'+''.join(fonts)+'</style></defs>',
         f'<rect width="{W}" height="{H}" fill="white"/>',f'<g transform="translate({DX} {DY})">']
    for i,s in enumerate(SCENE):
        label=escape(s['name'],{'"':'&quot;'})
        out.append(f'<g id="object-{i:04d}" data-name="{label}"><title>{escape(s["name"])}</title>')
        if s['kind'] in ('path','ellipse'):
            style=f'fill="{s["fill"] or "none"}" stroke="{s["stroke"] or "none"}" stroke-width="{s["width"]}" opacity="{s["opacity"]}"'
            if s['kind']=='path':
                d=' '.join(c[0]+' '+' '.join(f'{v:.4f}' for v in c[1:]) for c in s['commands'])
                out.append(f'<path d="{d}" {style}/>')
            else:
                out.append('<ellipse '+ ' '.join(f'{k}="{s[k]}"' for k in ('cx','cy','rx','ry'))+' '+style+'/>')
        elif s['kind']=='image':
            out.append(f'<image x="{s["x"]}" y="{s["y"]}" width="{s["w"]}" height="{s["h"]}" xlink:href="data:image/jpeg;base64,{s["data"]}"/>')
        elif s['kind']=='text':
            widths=[]
            for r in s['runs']:
                family=r.get('font',s['font'])
                font_path=(ROOT/'fonts'/('ComicNeue-Bold.ttf' if s['bold'] else 'ComicNeue-Regular.ttf')) if family=='Comic Neue' else ROOT/'fonts'/('DejaVuSerif-Bold.ttf' if s['bold'] else 'DejaVuSerif.ttf')
                f=ImageFont.truetype(str(font_path),round(s['size']*(.73 if r.get('sub') else 1)*100))
                widths.append(f.getlength(r['text'])/100)
            x=s['x']+((s['w']-sum(widths))/2 if s['align']=='center' else 0)
            y=s['y']+s['h']/2+s['size']*.32
            for r in s['runs']:
                sub=r.get('sub',False)
                attr=f'font-size="{s["size"]*(.73 if sub else 1)}"'
                if r.get('italic'):attr+=' font-style="italic"'
                attr+=f' font-family="{r.get("font",s["font"])}"'
                out.append(f'<text x="{x}" y="{y+(.24*s["size"] if sub else 0)}" {attr} font-weight="{700 if s["bold"] else 400}" fill="{s["fill"]}" xml:space="preserve">{escape(r["text"])}</text>')
                x+=widths.pop(0)
        out.append('</g>')
    out.extend(['</g>','</svg>'])
    return ''.join(out)


def el(tag, **attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items():e.set(k,str(v))
    return e


def paint(sh,s):
    if s['fill']:
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(s['fill'][1:])
        if s.get('opacity',1)<1:
            sh.fill._xPr.solidFill[0].append(el('a:alpha',val=round(s['opacity']*100000)))
    else:sh.fill.background()
    if s['stroke']:
        sh.line.color.rgb=RGBColor.from_string(s['stroke'][1:]);sh.line.width=Pt(s['width'])
        if s.get('opacity',1)<1:
            sh.line._get_or_add_ln().find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr').append(el('a:alpha',val=round(s['opacity']*100000)))
    else:sh.line.fill.background()


def ppt_render():
    prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
    prs.core_properties.title='AdvFD Figure 3 — editable reconstruction'
    prs.core_properties.subject='Reference figure from arXiv:2608.11205v1; native shapes and text'
    prs.core_properties.comments='Reference source retained. Editable scatter locations are illustrative redraws, not measured data.'
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.name='Editable Figure 3'
    # Group semantic components so dots and contours can move together.
    groups={}
    for s in SCENE:
        prefix=s['name'].split('/')[0]
        if prefix not in groups:
            groups[prefix]=slide.shapes.add_group_shape();groups[prefix].name=prefix
        target=groups[prefix].shapes
        if s['kind']=='ellipse':
            sh=target.add_shape(MSO_SHAPE.OVAL,Pt(s['cx']-s['rx']+DX),Pt(s['cy']-s['ry']+DY),Pt(2*s['rx']),Pt(2*s['ry']))
            paint(sh,s)
        elif s['kind']=='path':
            coords=[(c[i],c[i+1]) for c in s['commands'] for i in range(1,len(c),2)]
            x0=min(p[0] for p in coords);y0=min(p[1] for p in coords)
            x1=max(p[0] for p in coords);y1=max(p[1] for p in coords)
            w=max(x1-x0,.01);h=max(y1-y0,.01)
            sh=target.add_shape(MSO_SHAPE.RECTANGLE,Pt(x0+DX),Pt(y0+DY),Pt(w),Pt(h))
            spPr=sh._element.spPr
            spPr.remove(spPr.prstGeom)
            geom=el('a:custGeom')
            for tag in ['a:avLst','a:gdLst','a:ahLst','a:cxnLst']:geom.append(el(tag))
            geom.append(el('a:rect',l='0',t='0',r='r',b='b'))
            pl=el('a:pathLst');pa=el('a:path',w=round(w*1000),h=round(h*1000))
            for c in s['commands']:
                if c[0]=='Z':pa.append(el('a:close'));continue
                cmd=el({'M':'a:moveTo','L':'a:lnTo','C':'a:cubicBezTo'}[c[0]])
                for i in range(1,len(c),2):cmd.append(el('a:pt',x=round((c[i]-x0)*1000),y=round((c[i+1]-y0)*1000)))
                pa.append(cmd)
            pl.append(pa);geom.append(pl);spPr.insert(1,geom)
            paint(sh,s)
        elif s['kind']=='image':
            sh=target.add_picture(io.BytesIO(base64.b64decode(s['data'])),Pt(s['x']+DX),Pt(s['y']+DY),Pt(s['w']),Pt(s['h']))
        else:
            sh=target.add_textbox(Pt(s['x']+DX),Pt(s['y']+DY),Pt(s['w']),Pt(s['h']))
            tf=sh.text_frame;tf.clear();tf.word_wrap=False
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            tf.vertical_anchor=MSO_ANCHOR.MIDDLE
            p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER if s['align']=='center' else PP_ALIGN.LEFT
            p.space_before=p.space_after=Pt(0)
            for r in s['runs']:
                run=p.add_run();run.text=r['text'];f=run.font
                f.name=r.get('font',s['font']);f.size=Pt(s['size']*(.73 if r.get('sub') else 1))
                f.bold=s['bold'];f.italic=r.get('italic',False)
                f.color.rgb=RGBColor.from_string(s['fill'][1:])
                if r.get('sub'):run._r.get_or_add_rPr().set('baseline','-25000')
        sh.name=s['name']
    slide.notes_slide.notes_text_frame.text=(
        'Editable reconstruction of AdvFD (Gao et al., 2026), arXiv:2608.11205v1, Figure 3. '
        'Original module/edge geometry retained from the authors’ vector PDF. '
        'Text/formulas, icons, contours and scatter points are native editable objects. '
        'The two noise tiles are replaceable raster pictures. Scatter positions were redrawn and are illustrative. '
        'Use Comic Neue (supplied fonts) and DejaVu Serif. Ungroup or double-click a group to edit its contents. '
        'This page reproduces AdvFD, not the local classifier-guidance method. '
        'The original D-step static branch is retained for visual fidelity; its objective only optimizes the adverse branch.')
    reference=prs.slides.add_slide(prs.slide_layouts[6]);reference.name='Original reference (not editable)'
    pix=PAGE.get_pixmap(matrix=fitz.Matrix(2.5,2.5),alpha=False)
    pix.save(ROOT/'reference'/'advfd_figure3_original.png')
    reference.shapes.add_picture(io.BytesIO(pix.tobytes('png')),Pt(DX),Pt(DY),Pt(821),Pt(464))
    reference.notes_slide.notes_text_frame.text='Original Figure 3 from https://arxiv.org/abs/2608.11205v1, for visual comparison. Edit slide 1.'
    prs.save(ROOT/'advfd_figure3_editable.pptx')


if __name__=='__main__':
    build_scene()
    svg=svg_render()
    (ROOT/'advfd_figure3_editable.svg').write_text(svg)
    (ROOT/'scene.json').write_text(json.dumps(SCENE,ensure_ascii=False,indent=2))
    cairosvg.svg2png(bytestring=svg.encode(),write_to=str(ROOT/'advfd_figure3_preview.png'),output_width=2000,output_height=round(2000*H/W))
    cairosvg.svg2pdf(bytestring=svg.encode(),write_to=str(ROOT/'advfd_figure3_editable.pdf'))
    ppt_render()
    print(f'Created {len(SCENE)} editable scene objects in {ROOT}')
